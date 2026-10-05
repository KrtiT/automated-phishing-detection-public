from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = (
    ROOT
    / "deliverables"
    / "Tallam_Praxis_Advisor_Update_2026-08-20_With_Speaker_Notes.pptx"
)
OUTPUT = (
    ROOT
    / "deliverables"
    / "Tallam_Praxis_Advisor_Update_2026-09-03_With_Speaker_Notes.pptx"
)

NAVY = RGBColor(0x05, 0x3C, 0x5A)
TEAL = RGBColor(0x00, 0x78, 0x9A)
TEXT = RGBColor(0x2F, 0x33, 0x37)
MUTED = RGBColor(0x5F, 0x69, 0x72)
LIGHT = RGBColor(0xF4, 0xF7, 0xF9)
LIGHT_BLUE = RGBColor(0xE9, 0xF3, 0xF7)
LINE = RGBColor(0xD6, 0xDF, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run_font(
    run,
    size: float,
    *,
    bold: bool = False,
    color: RGBColor = TEXT,
    name: str = "Arial",
) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def delete_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def clear_slide_body(slide) -> None:
    title = slide.shapes.title
    for shape in list(slide.shapes):
        if not shape.is_placeholder:
            delete_shape(shape)
            continue
        if shape is title:
            continue
        placeholder_type = shape.placeholder_format.type
        if placeholder_type == PP_PLACEHOLDER.SLIDE_NUMBER:
            continue
        if shape.has_text_frame:
            shape.text_frame.clear()


def set_title(slide, text: str) -> None:
    title = slide.shapes.title
    if title is None:
        raise ValueError("expected a title placeholder")
    title.text_frame.clear()
    paragraph = title.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.LEFT
    set_run_font(paragraph.runs[0], 23, bold=True)


def add_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    size: float,
    bold: bool = False,
    color: RGBColor = TEXT,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    font_name: str = "Arial",
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> object:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    set_run_font(
        paragraph.runs[0], size, bold=bold, color=color, name=font_name
    )
    return box


def add_panel(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    heading: str,
    items: list[str],
    *,
    font_size: float = 14.0,
    fill: RGBColor = LIGHT,
    heading_color: RGBColor = NAVY,
    spacing: float = 6.0,
) -> None:
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = fill
    panel.line.color.rgb = LINE
    panel.line.width = Pt(0.8)
    panel.adjustments[0] = 0.08

    add_text(
        slide,
        x + 0.18,
        y + 0.10,
        w - 0.36,
        0.32,
        heading,
        size=13,
        bold=True,
        color=heading_color,
    )
    body = slide.shapes.add_textbox(
        Inches(x + 0.20), Inches(y + 0.48), Inches(w - 0.40), Inches(h - 0.58)
    )
    frame = body.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"\u2022 {item}"
        paragraph.space_after = Pt(spacing)
        paragraph.line_spacing = 1.02
        set_run_font(paragraph.runs[0], font_size)


def add_status_row(
    slide,
    y: float,
    label: str,
    statement: str,
    status: str,
    *,
    height: float = 1.18,
    statement_size: float = 12.2,
    status_size: float = 10.8,
) -> None:
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.48),
        Inches(y),
        Inches(9.02),
        Inches(height),
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = LIGHT
    panel.line.color.rgb = LINE
    panel.line.width = Pt(0.8)
    panel.adjustments[0] = 0.06

    add_text(
        slide,
        0.68,
        y + 0.12,
        0.65,
        0.34,
        label,
        size=13,
        bold=True,
        color=NAVY,
    )
    add_text(
        slide,
        1.33,
        y + 0.10,
        6.26,
        height - 0.20,
        statement,
        size=statement_size,
        valign=MSO_ANCHOR.MIDDLE,
    )

    status_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.82),
        Inches(y + 0.15),
        Inches(1.43),
        Inches(height - 0.30),
    )
    status_box.fill.solid()
    status_box.fill.fore_color.rgb = LIGHT_BLUE
    status_box.line.color.rgb = TEAL
    status_box.line.width = Pt(0.8)
    status_box.adjustments[0] = 0.12
    frame = status_box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.text = status
    paragraph.alignment = PP_ALIGN.CENTER
    set_run_font(paragraph.runs[0], status_size, bold=True, color=NAVY)


def add_note(
    slide,
    text: str,
    *,
    y: float = 5.48,
    height: float = 0.44,
    font_size: float = 11.1,
) -> None:
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.48),
        Inches(y),
        Inches(9.02),
        Inches(height),
    )
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_BLUE
    box.line.color.rgb = LIGHT_BLUE
    box.adjustments[0] = 0.08
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Inches(0.16)
    frame.margin_right = Inches(0.16)
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    set_run_font(paragraph.runs[0], font_size, bold=True, color=NAVY)


def add_metric_card(
    slide, x: float, label: str, value: str, qualifier: str = ""
) -> None:
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(1.10),
        Inches(2.12),
        Inches(1.18),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT
    card.line.color.rgb = LINE
    card.line.width = Pt(0.8)
    card.adjustments[0] = 0.08
    add_text(slide, x + 0.16, 1.22, 1.80, 0.22, label.upper(), size=9.0, bold=True, color=TEAL)
    add_text(slide, x + 0.16, 1.49, 1.80, 0.37, value, size=20.0, bold=True, color=NAVY)
    if qualifier:
        add_text(slide, x + 0.16, 1.90, 1.80, 0.20, qualifier, size=8.8, color=MUTED)


def add_table_panel(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    heading: str,
    columns: list[tuple[float, str]],
    rows: list[list[str]],
    *,
    first_col_bold: bool = True,
) -> None:
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = LIGHT
    panel.line.color.rgb = LINE
    panel.line.width = Pt(0.8)
    panel.adjustments[0] = 0.06
    add_text(slide, x + 0.18, y + 0.11, w - 0.36, 0.30, heading, size=13, bold=True, color=NAVY)

    inner_x = x + 0.18
    table_w = w - 0.36
    header_y = y + 0.54
    row_h = (h - 0.74) / (len(rows) + 1)
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(inner_x),
        Inches(header_y),
        Inches(table_w),
        Inches(row_h),
    )
    header.fill.solid()
    header.fill.fore_color.rgb = LIGHT_BLUE
    header.line.color.rgb = LIGHT_BLUE

    col_x = inner_x
    for fraction, label in columns:
        add_text(
            slide,
            col_x + 0.07,
            header_y + 0.04,
            table_w * fraction - 0.14,
            row_h - 0.08,
            label,
            size=9.2,
            bold=True,
            color=TEAL,
            valign=MSO_ANCHOR.MIDDLE,
        )
        col_x += table_w * fraction

    for row_index, row in enumerate(rows):
        row_y = header_y + row_h * (row_index + 1)
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(inner_x),
            Inches(row_y),
            Inches(table_w),
            Inches(row_h),
        )
        background.fill.solid()
        background.fill.fore_color.rgb = WHITE if row_index % 2 == 0 else LIGHT
        background.line.color.rgb = LINE
        background.line.width = Pt(0.4)
        col_x = inner_x
        for col_index, ((fraction, _), value) in enumerate(zip(columns, row, strict=True)):
            add_text(
                slide,
                col_x + 0.07,
                row_y + 0.03,
                table_w * fraction - 0.14,
                row_h - 0.06,
                value,
                size=10.2,
                bold=first_col_bold and col_index == 0,
                color=NAVY if first_col_bold and col_index == 0 else TEXT,
                valign=MSO_ANCHOR.MIDDLE,
            )
            col_x += table_w * fraction


def build_title_slide(slide) -> None:
    frame = slide.shapes[0].text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = Inches(0.16)
    frame.margin_top = 0
    frame.margin_bottom = 0
    entries = [
        ("Praxis Title: ", "Automated Phishing Detection for Frontier AI Inference", 17.2),
        ("Update Focus: ", "Current Research Implementation and Evidence Controls", 17.2),
        ("Name and Date: ", "Krti Tallam | September 3, 2026", 17.2),
        (
            "Elevator Pitch: ",
            "Since August 20, I moved from a plan to a reproducible development-data milestone and tightened the evidence chain behind labels, splits, and later hypothesis tests.",
            15.2,
        ),
    ]
    for index, (label, value, size) in enumerate(entries):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(15 if index < 3 else 0)
        paragraph.line_spacing = 1.04
        label_run = paragraph.add_run()
        label_run.text = label
        set_run_font(label_run, size, bold=True, color=WHITE)
        value_run = paragraph.add_run()
        value_run.text = value
        set_run_font(value_run, size, color=WHITE)


def build_feedback_slide(slide) -> None:
    set_title(slide, "Response to August 20 Advisor Feedback")
    add_panel(
        slide,
        0.48,
        1.08,
        4.35,
        3.92,
        "What I heard",
        [
            "Research outcomes must not depend on any individual's judgment.",
            "A reader should be able to trace who supplied the reference classifications and how they become study outcomes.",
            "Ambiguous, invalid, or contradictory cases need one consistent handling rule.",
        ],
        font_size=13.7,
    )
    add_panel(
        slide,
        5.17,
        1.08,
        4.35,
        3.92,
        "Implemented response",
        [
            "Publisher-provided reference classifications are the only input to the deterministic 0/1 study mapping.",
            "No researcher assigns, corrects, or overrides an outcome; the same source record always follows the same rule.",
            "Invalid or conflicting cases are mechanically quarantined, and the source limitation is disclosed.",
        ],
        font_size=13.5,
    )
    add_note(
        slide,
        "Repeatability means the same inputs yield the same decision; it does not make publisher labels infallible.",
        y=5.24,
    )


def build_progress_slide(slide) -> None:
    set_title(slide, "Work Completed Since August 20")
    add_panel(
        slide,
        0.48,
        1.08,
        4.35,
        3.86,
        "Public, reviewable milestone",
        [
            "Branch: github.com/KrtiT/automated-phishing-detection-public/tree/praxis-realignment-v3",
            "Milestone commits: 9b7115a, a17291f, and 9d47009.",
            "GitHub CI run 33810012730 passed against the published 175-test baseline.",
        ],
        font_size=12.1,
        spacing=7,
    )
    add_panel(
        slide,
        5.17,
        1.08,
        4.35,
        3.86,
        "What the milestone establishes",
        [
            "The prior v2 state remains preserved at tag a5eceec.",
            "Preparation now runs end to end: source validation, deterministic mapping, quarantine, domain grouping, split allocation, and aggregate evidence.",
            "Model fitting and hypothesis tests remain ahead; source-release provenance packaging is active work, not a published claim.",
        ],
        font_size=12.9,
    )
    add_note(
        slide,
        "Completed: development-data preparation. Still ahead: model, routing-outcome, HTTP, and external-evaluation evidence.",
        y=5.18,
    )


def build_label_chain_slide(slide) -> None:
    set_title(slide, "Label and Provenance Chain")
    add_status_row(
        slide,
        1.08,
        "1",
        "The research observations are actual URL strings published in UCI PhiUSIIL; they were not generated for this study.",
        "Observed source",
        height=1.20,
        statement_size=12.4,
    )
    add_status_row(
        slide,
        2.45,
        "2",
        "PhiUSIIL uses native 0 = phishing and 1 = legitimate. Its paper attributes legitimate cases to Open PageRank and phishing cases to PhishTank, OpenPhish, and MalwareWorld. The study flips the encoding only.",
        "Mechanical map",
        height=1.32,
        statement_size=11.5,
    )
    add_status_row(
        slide,
        3.94,
        "3",
        "Publisher labels can contain measurement error. No case is personally relabeled; invalid records and conflicting domain groups are quarantined by rule.",
        "Validity control",
        height=1.24,
        statement_size=12.0,
    )
    add_note(
        slide,
        "The chain is inspectable and repeatable; its stated source limitation keeps repeatability distinct from infallibility.",
        y=5.42,
    )


def build_snapshot_slide(slide) -> None:
    set_title(slide, "Preparation Run Snapshot")
    for x, label, value, qualifier in [
        (0.48, "Input", "235,795", "publisher rows"),
        (2.78, "Retained", "233,536", "99.04% of input"),
        (5.08, "Quarantined", "2,259", "0.96% of input"),
        (7.38, "Domains", "197,105", "registrable groups"),
    ]:
        add_metric_card(slide, x, label, value, qualifier)

    add_table_panel(
        slide,
        0.48,
        2.52,
        5.62,
        2.56,
        "Group-disjoint retained split",
        [(0.42, "Partition"), (0.29, "Rows"), (0.29, "Domains")],
        [
            ["Train", "166,248", "137,973"],
            ["Validation", "32,695", "29,566"],
            ["Group-test", "34,593", "29,566"],
        ],
    )
    add_table_panel(
        slide,
        6.30,
        2.52,
        3.22,
        2.56,
        "Quarantine reasons",
        [(0.70, "Rule"), (0.30, "Rows")],
        [
            ["Invalid URL", "1,380"],
            ["Same-label duplicate", "877"],
            ["Conflicting group", "2"],
        ],
    )
    add_note(
        slide,
        "These counts are preparation evidence about data handling and partitioning; they are not model results.",
        y=5.30,
    )


def build_questions_slide(slide) -> None:
    set_title(slide, "Active Research Questions")
    add_note(
        slide,
        "Preparation is complete and future-only routing mechanics are implemented; H1, H2, and H3 remain undecided.",
        y=1.00,
        height=0.58,
        font_size=11.0,
    )
    add_status_row(
        slide,
        1.77,
        "RQ1",
        "What incremental value do structural URL features and character-level representations provide under registrable-domain-disjoint and external evaluation?",
        "Active question",
        height=1.10,
        statement_size=11.7,
    )
    add_status_row(
        slide,
        3.02,
        "RQ2",
        "Can GMM-based monitoring detect an external source/domain shift and guide escalation without exceeding the low-FPR operating constraint?",
        "Active question",
        height=1.10,
        statement_size=11.7,
    )
    add_status_row(
        slide,
        4.27,
        "RQ3",
        "What detection, escalation, throughput, and latency tradeoffs determine whether the fixed cascade is viable inline?",
        "Active question",
        height=1.10,
        statement_size=11.7,
    )
    add_note(
        slide,
        "Evidence status: no confirmatory hypothesis result has been run; the questions retain their matrix v1.2 wording.",
        y=5.52,
    )


def build_decision_rules_slide(slide) -> None:
    set_title(slide, "Decision Rules and Target Basis")
    add_status_row(
        slide,
        1.08,
        "H1",
        "At <= 1% observed FPR on both primary test sets, both domain-clustered 95% bootstrap lower bounds must be > 0: Logistic-L1 over length-only recall, and cascade over Logistic-L1 recall.",
        "All gates",
        height=1.26,
        statement_size=11.0,
    )
    add_status_row(
        slide,
        2.50,
        "H2",
        ">= 80% shift-window detection; <= 5% independent reference false alerts; routed-policy FPR <= 1%; and a 95% lower bound for external false-negative-rate reduction > 0.",
        "All gates",
        height=1.20,
        statement_size=11.1,
    )
    add_status_row(
        slide,
        3.86,
        "H3",
        "Each system: certified-registry FPR <= 1% and Tranco control alert rate <= 1%; recall lower bound >= -0.02; invocation <= 30%; real-HTTP p95 <= 200 ms at concurrency 64; errors < 0.1%.",
        "All gates",
        height=1.28,
        statement_size=10.8,
    )
    add_note(
        slide,
        "These are study-defined operating gates, not literature-prescribed or achieved values. A pass requires every gate, not one high score.",
        y=5.34,
        height=0.58,
        font_size=10.7,
    )


def build_novelty_slide(slide) -> None:
    set_title(slide, "Novelty After the Current Literature Check")
    add_status_row(
        slide,
        1.02,
        "1",
        "Ahamed et al. (2026): an integrated protocol for adversarial robustness, generalization, and explanation stability in URL phishing detection.",
        "Evaluation",
        height=0.98,
        statement_size=11.3,
    )
    add_status_row(
        slide,
        2.14,
        "2",
        "ExpertFusion (Hussain et al., 2026): calibrated multi-expert URL decision fusion under distribution shift and target-prior uncertainty.",
        "Fusion",
        height=0.98,
        statement_size=11.3,
    )
    add_status_row(
        slide,
        3.26,
        "3",
        "Alajaji (2026): a classical-first selective cascade for resource-constrained phishing email detection.",
        "Cascade",
        height=0.98,
        statement_size=11.5,
    )
    add_panel(
        slide,
        0.48,
        4.42,
        9.02,
        1.08,
        "Bounded contribution",
        [
            "Prospective joint systems evaluation of a frozen structural-to-character-transformer URL cascade, GMM-triggered future-only routing, external source/domain shift, and jointly constrained FPR, compute, real-HTTP tail latency, and reliability."
        ],
        font_size=11.3,
        fill=LIGHT_BLUE,
        heading_color=TEAL,
        spacing=0,
    )
    add_note(
        slide,
        "The components are established; the contribution is their predeclared joint evaluation, with no component-level priority claim.",
        y=5.68,
        height=0.34,
        font_size=9.8,
    )


def build_reproducibility_slide(slide) -> None:
    set_title(slide, "Reproducibility and Source Freeze")
    add_panel(
        slide,
        0.48,
        1.08,
        4.35,
        3.92,
        "Pinned in the current repository",
        [
            "UCI archive SHA-256 prefix: 0a639fd0...",
            "Extracted CSV SHA-256 prefix: a236549c...",
            "Public Suffix List commit: 0f1fa47...",
            "The uv lock, split/output hashes, and aggregate preparation record complete the current verification chain.",
        ],
        font_size=12.8,
    )
    add_panel(
        slide,
        5.17,
        1.08,
        4.35,
        3.92,
        "Today's source-release packaging",
        [
            "Package the exact CC BY 4.0 UCI ZIP with its source record and checksums outside Git history.",
            "Keep the repository's identity and verification path compact while preserving the exact source bytes for public review.",
            "Describe this packaging as active work; it has not yet been published as a release.",
            "PhishVN remains untouched in this milestone.",
        ],
        font_size=12.5,
    )
    add_note(
        slide,
        "Freeze chain: source bytes -> extraction -> suffix rules -> deterministic split -> aggregate evidence.",
        y=5.24,
    )


def build_next_work_slide(slide) -> None:
    set_title(slide, "Next Executable Work")
    add_panel(
        slide,
        0.48,
        1.08,
        4.35,
        3.64,
        "Active sequence",
        [
            "Complete the source-release and public provenance record.",
            "Implement frozen URL features plus length-only and Logistic-L1 baselines using train and validation only.",
            "Lock thresholds and the cascade band, then run the untouched group-test once.",
            "Continue transformer, GMM, and real-HTTP milestones; then perform the frozen external evaluation.",
        ],
        font_size=12.2,
        spacing=5.5,
    )
    add_panel(
        slide,
        5.17,
        1.08,
        4.35,
        3.64,
        "Exact review locations",
        [
            "README",
            "data/sources.json",
            "src/automated_phishing_detection/phiusiil.py",
            "tests/ and the reports summary",
            "research-alignment protocol matrix v1.2",
        ],
        font_size=12.0,
        spacing=5,
    )
    add_panel(
        slide,
        0.48,
        4.86,
        9.02,
        0.90,
        "Execution rule",
        [
            "Each artifact is frozen before it controls later evidence: development choices use train/validation, group-test is scored once, and external evaluation follows the remaining system milestones."
        ],
        font_size=10.8,
        fill=LIGHT_BLUE,
        heading_color=TEAL,
        spacing=0,
    )


def build_appendix_slide(slide) -> None:
    set_title(slide, "Appendix: Code and Data Evidence Snapshot")
    add_table_panel(
        slide,
        0.48,
        1.08,
        5.02,
        3.92,
        "Aggregate retained split",
        [(0.40, "Partition"), (0.30, "Rows"), (0.30, "Domains")],
        [
            ["Train", "166,248", "137,973"],
            ["Validation", "32,695", "29,566"],
            ["Group-test", "34,593", "29,566"],
            ["Total", "233,536", "197,105"],
        ],
    )

    terminal = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.70),
        Inches(1.08),
        Inches(3.82),
        Inches(2.18),
    )
    terminal.fill.solid()
    terminal.fill.fore_color.rgb = NAVY
    terminal.line.color.rgb = TEAL
    terminal.line.width = Pt(0.8)
    terminal.adjustments[0] = 0.06
    add_text(slide, 5.92, 1.24, 3.38, 0.27, "Published checks", size=12.6, bold=True, color=WHITE)
    lines = [
        "$ pytest",
        "  175 passed",
        "$ gh run view 33810012730",
        "  conclusion: success",
    ]
    box = slide.shapes.add_textbox(Inches(5.92), Inches(1.62), Inches(3.35), Inches(1.36))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.space_after = Pt(3)
        set_run_font(paragraph.runs[0], 10.1, color=WHITE, name="Courier New")

    milestone = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.70),
        Inches(3.46),
        Inches(3.82),
        Inches(1.54),
    )
    milestone.fill.solid()
    milestone.fill.fore_color.rgb = LIGHT
    milestone.line.color.rgb = LINE
    milestone.line.width = Pt(0.8)
    milestone.adjustments[0] = 0.06
    add_text(slide, 5.92, 3.61, 3.38, 0.28, "Public milestone chain", size=12.4, bold=True, color=NAVY)
    add_text(slide, 5.92, 4.01, 3.38, 0.33, "9b7115a -> a17291f -> 9d47009", size=10.2, color=TEXT, font_name="Courier New")
    add_text(slide, 5.92, 4.46, 3.38, 0.24, "Preserved v2 tag: a5eceec", size=10.0, bold=True, color=TEAL)

    add_text(
        slide,
        0.58,
        5.18,
        8.82,
        0.30,
        "https://github.com/KrtiT/automated-phishing-detection-public/tree/praxis-realignment-v3",
        size=10.2,
        bold=True,
        color=TEAL,
        align=PP_ALIGN.CENTER,
    )
    add_note(
        slide,
        "Reserve slide for inspection on request. It contains aggregate evidence only and no row-level URL observations.",
        y=5.58,
        height=0.38,
        font_size=10.1,
    )


SLIDE_NOTES = [
    [
        "Since the August 20 discussion, I have moved from a study plan to a reproducible development-data milestone. The concrete observation is an end-to-end preparation run with a public branch, pinned inputs, fixed rules, aggregate outputs, and passing checks.",
        "My interpretation is deliberately narrower than a research result: the milestone strengthens the evidence chain that later model comparisons will rely on. It does not decide H1, H2, or H3.",
        "The observations are actual URL strings from the UCI PhiUSIIL release. They are not generated research records, and no row-level URL is shown in this deck.",
        "Today I want to clarify the boundary between completed preparation evidence and the model, routing, HTTP, and external evidence that comes next.",
        "Evidence map: matrix v1.2 development-data role, validity controls, research safeguards, and decision matrix.",
    ],
    [
        "The August 20 concern was that an outcome should not depend on one person's judgment. I translated that concern into an auditable procedure rather than a case-by-case decision process.",
        "The observed implementation uses publisher-provided reference classifications, a deterministic binary mapping, and mechanical quarantine. No researcher assigns, corrects, or overrides a record's outcome.",
        "Two people applying the same source record and versioned rule should reach the same study outcome. Someone can still reasonably question the publisher's classification or the construct itself; repeatability does not remove measurement error.",
        "Labels can be wrong. The defensible response is to disclose that source limitation, preserve provenance, avoid personal relabeling, and quarantine records that fail the stated validity rules.",
        "My interpretation is that this removes individual discretion from the outcome procedure while keeping the publisher-label limitation visible.",
        "Today I want to clarify whether this repeatability-versus-infallibility explanation answers the underlying concern directly enough for the methodology narrative.",
        "Evidence map: matrix v1.2 outcome-label procedure, validity controls, and research safeguards.",
    ],
    [
        "The observable work since August 20 is public: the praxis-realignment-v3 branch, milestone commits 9b7115a, a17291f, and 9d47009, and GitHub Actions run 33810012730 with a successful conclusion.",
        "The published baseline contains 175 passing tests. The earlier v2 state remains identifiable at tag a5eceec, so the prospective work does not rewrite that reference point.",
        "The preparation pipeline now executes source checks, mapping, quarantine, registrable-domain grouping, deterministic partitioning, and aggregate reporting end to end.",
        "My interpretation is that this is a reviewable data-preparation milestone. It is not evidence that any classifier, cascade, or monitoring hypothesis has succeeded.",
        "Today's source-release packaging and provenance record remain active work. I am not presenting uncommitted provenance changes or a release as already public.",
        "Today I want to clarify which review surface is most useful first: the README and source record, the pipeline implementation and tests, or the aggregate report.",
        "Evidence map: public branch, CI run, report summary, and matrix v1.2 safeguards.",
    ],
    [
        "The source observations are the URL strings distributed in UCI PhiUSIIL. The study consumes those observed strings; it does not create substitute research observations.",
        "PhiUSIIL's native convention is 0 for phishing and 1 for legitimate. The study reverses that encoding so the local positive class is phishing, but it does not change the publisher's classification.",
        "The accompanying paper describes Open PageRank as the legitimate source and PhishTank, OpenPhish, and MalwareWorld as the phishing sources. That is source provenance, not a personal labeling exercise.",
        "If a URL is invalid or a registrable-domain group contains conflicting labels, the pipeline follows a fixed quarantine rule. I do not adjudicate which label feels more plausible.",
        "Publisher labels can still contain measurement error, stale classifications, or source-specific bias. My interpretation is that the procedure is repeatable and transparent, not infallible.",
        "Today I want to clarify the most concise wording for that limitation so the reader understands both the control and its boundary.",
        "Evidence map: matrix v1.2 development data, outcome mapping, and validity controls; implementation in phiusiil.py and its tests.",
    ],
    [
        "The observed preparation run begins with 235,795 publisher rows. It retains 233,536 rows and quarantines 2,259, leaving 197,105 registrable-domain groups.",
        "The retained data are partitioned by domain group: 166,248 train rows across 137,973 domains, 32,695 validation rows across 29,566 domains, and 34,593 group-test rows across 29,566 domains.",
        "The quarantine ledger contains 1,380 invalid URLs, 877 same-label duplicates, and two conflicting groups. Those three counts reconcile exactly to the 2,259 quarantined rows.",
        "The interpretation is limited to data handling: the pipeline applied the declared mapping, grouping, partition, and quarantine controls and produced a reconcilable aggregate record.",
        "These numbers say nothing about recall, false-positive rate, calibration, latency, or any hypothesis result. The untouched group-test has not been used to select a model or threshold.",
        "Today I want to clarify whether this aggregate view makes the preparation-versus-performance distinction sufficiently explicit.",
        "Evidence map: reports preparation summary, split manifest hashes, and matrix v1.2 development-data safeguards.",
    ],
    [
        "These are the exact active questions in public matrix v1.2. RQ1 isolates the incremental value of structural and character representations under group-disjoint and external evaluation.",
        "RQ2 asks whether GMM monitoring detects external source or domain shift and whether future-only escalation improves errors while respecting the low-FPR constraint.",
        "RQ3 asks whether the frozen cascade is viable when detection, escalation, throughput, latency, and reliability are considered together.",
        "The observed status is preparation complete and future-only routing mechanics implemented. The confirmatory evidence for H1, H2, and H3 has not been run, so each hypothesis remains undecided.",
        "My interpretation is that implemented mechanics make the questions executable; mechanics alone do not answer them.",
        "Today I want to clarify that this status language cleanly separates active questions, implemented controls, and evidence not yet run.",
        "Evidence map: matrix v1.2 RQ1 through RQ3 and the preparation/routing implementation tests.",
    ],
    [
        "The exact numeric gates are study-defined operating rules. The literature motivates low false-positive operation, shift evaluation, selective computation, and latency measurement, but it does not prescribe these exact numbers or establish that this study has achieved them.",
        "For H1, each comparison must remain at or below one percent observed FPR on both primary test sets, and both domain-clustered bootstrap lower bounds for the planned recall improvements must be above zero.",
        "For H2, the GMM policy must meet all four conditions: at least 80 percent window detection, no more than five percent independent false alerts, routed-policy FPR at or below one percent, and a positive lower bound for false-negative-rate reduction.",
        "For H3, both the fixed cascade and transformer-only system face the dual one-percent safeguards: certified-registry FPR and Tranco control alert rate. The recall lower bound must be at least minus 0.02, invocation no more than 30 percent, real-HTTP p95 no more than 200 milliseconds at concurrency 64, and errors below 0.1 percent.",
        "The observed status is that none of these gates has been evaluated. My interpretation is that requiring every gate prevents one high score from masking unsafe false positives, excessive compute, poor tail latency, or request failures.",
        "Today I want to clarify how best to state the target basis: explicit design choices for a stringent joint operating test, reported separately from literature findings.",
        "Evidence map: matrix v1.2 H1, H2, H3, complete decision rules, and target-basis note.",
    ],
    [
        "The observed literature evidence confirms that the individual components are established. Ahamed and colleagues provide an integrated URL evaluation protocol spanning adversarial robustness, generalization, and explanation stability.",
        "ExpertFusion evaluates calibrated multi-expert URL decisions under distribution shift and unknown target prevalence, including registered-domain-aware evaluation. Alajaji evaluates a classical-first selective cascade, although the task is phishing email detection rather than this URL-only system.",
        "Those papers narrow the novelty claim. I am not claiming the first cascade, the first transformer URL detector, the first shift detector, or a new component in isolation.",
        "My interpretation of the remaining contribution is a prospective joint systems evaluation: a frozen structural-to-character-transformer URL cascade, GMM-triggered future-only routing, external source/domain shift, and simultaneous FPR, compute, real-HTTP tail-latency, and reliability gates.",
        "This is a bounded synthesis from the current search, not proof that no adjacent paper exists. The contribution will stand or fall on the predeclared joint evidence, not a priority claim.",
        "Today I want to clarify whether that systems-evaluation contribution is stated narrowly enough while remaining doctoral in scope.",
        "Evidence map: current Chapter 2 literature notes and matrix v1.2 problem, thesis, continuity, and decision rules.",
    ],
    [
        "The current repository records the UCI archive digest beginning 0a639fd0, the extracted CSV digest beginning a236549c, and Public Suffix List commit beginning 0f1fa47. The full values live in the source record and aggregate report.",
        "The uv lock, deterministic split and output hashes, and aggregate preparation record extend that chain from source identity through transformation outputs.",
        "The observable state today is a pinned repository and reproducible preparation evidence. The exact CC BY 4.0 UCI ZIP, source record, and checksums are now being packaged as the source release outside Git history.",
        "My interpretation is that this keeps immutable source bytes reviewable without using normal source-control history as a bulk-data store. I will describe the packaging as active work until the release artifact is actually public.",
        "PhishVN remains untouched in this milestone. There is no external row access, prediction, or outcome to report here.",
        "Today I want to clarify whether the slide's short hash prefixes and the repository review path give the right presentation-level detail, with full values left in the artifacts.",
        "Evidence map: data/sources.json, report summary, uv lock, output hashes, and matrix v1.2 research safeguards.",
    ],
    [
        "The next work is an executable sequence rather than a request to wait. First I will complete the source-release and public provenance record so the exact input identity is independently checkable.",
        "Second, I will implement the frozen URL feature set and the length-only and Logistic-L1 baselines. Training, feature decisions, and threshold selection will use only train and validation data.",
        "Third, I will lock thresholds and the cascade band. Fourth, I will score the untouched PhiUSIIL group-test once and preserve its prediction and summary artifacts.",
        "Fifth, I will continue the character transformer, GMM future-only routing, and real-HTTP milestones, followed by the frozen external evaluation. That is where the remaining H1, H2, and H3 evidence is produced.",
        "The exact review trail is the README, data/sources.json, phiusiil.py, tests, reports summary, and protocol matrix v1.2. Each location answers a different question: intent, source identity, transformation, behavior, aggregate evidence, and decision rule.",
        "The observed state is preparation complete with source-release packaging active. My interpretation is that this sequence protects each later result from development choices made after seeing evaluation evidence.",
        "Today I want to clarify which of those artifacts should anchor the next advisor walkthrough and which evidence summary would be most useful at the following update.",
        "Evidence map: matrix v1.2 proposed method and evidence, validity controls, research safeguards, and decision rules.",
    ],
    [
        "This is a reserve slide if Professor Etemadi asks to inspect the work. It contains the aggregate split, a concise test and CI excerpt, the public branch, milestone commit chain, and preserved v2 tag.",
        "The split totals reconcile to 233,536 retained rows and 197,105 registrable domains. No row-level URL appears here, so inspection remains focused on aggregate evidence and the public audit path.",
        "The terminal-style excerpt is a text rendering of the published evidence, not a simulated screenshot: the baseline reports 175 passed, and GitHub Actions run 33810012730 reports success.",
        "The test suite includes small unit-test fixtures to exercise parser and rule behavior. Those fixtures are not research observations and do not contribute to any preparation count or outcome.",
        "My interpretation is that the public chain is sufficient to reproduce and challenge the preparation milestone while keeping the untouched evaluation boundaries visible.",
        "Today I want to clarify whether he would prefer a code-first walkthrough, a source-and-hash walkthrough, or an aggregate-report walkthrough when this reserve slide is used.",
        "Evidence map: public branch and commits, CI run, reports summary, split manifest, and matrix v1.2 safeguards.",
    ],
]

SLIDE_TITLES = [
    "Current Research Implementation and Evidence Controls",
    "Response to August 20 Advisor Feedback",
    "Work Completed Since August 20",
    "Label and Provenance Chain",
    "Preparation Run Snapshot",
    "Active Research Questions",
    "Decision Rules and Target Basis",
    "Novelty After the Current Literature Check",
    "Reproducibility and Source Freeze",
    "Next Executable Work",
    "Appendix: Code and Data Evidence Snapshot",
]


def add_speaker_notes(slide, lines: list[str]) -> None:
    frame = slide.notes_slide.notes_text_frame
    frame.clear()
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"- {line}"
        paragraph.space_after = Pt(5)
        for run in paragraph.runs:
            set_run_font(run, 12)


def sanitize_inherited_package_metadata(path: Path) -> None:
    temp_path = path.with_name(f".{path.stem}.package-rewrite.pptx")
    app_ns = {
        "ep": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties",
        "vt": "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes",
    }
    notes_ns = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }
    with ZipFile(path) as source, ZipFile(
        temp_path, "w", compression=ZIP_DEFLATED
    ) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "docProps/app.xml":
                root = etree.fromstring(data)
                title_nodes = root.xpath(
                    "./ep:TitlesOfParts/vt:vector/vt:lpstr", namespaces=app_ns
                )
                if len(title_nodes) < len(SLIDE_TITLES):
                    raise ValueError("unexpected PowerPoint title cache")
                for node, title in zip(
                    title_nodes[-len(SLIDE_TITLES) :], SLIDE_TITLES, strict=True
                ):
                    node.text = title
                data = etree.tostring(
                    root,
                    xml_declaration=True,
                    encoding="UTF-8",
                    standalone=True,
                )
            elif entry.filename == "ppt/notesMasters/notesMaster1.xml":
                root = etree.fromstring(data)
                for node in root.xpath(
                    ".//a:fld[@type='datetimeFigureOut']/a:t", namespaces=notes_ns
                ):
                    node.text = "9/3/26"
                data = etree.tostring(
                    root,
                    xml_declaration=True,
                    encoding="UTF-8",
                    standalone=True,
                )
            target.writestr(entry, data)
    temp_path.replace(path)


def build_deck() -> Path:
    presentation = Presentation(TEMPLATE)
    if len(presentation.slides) != 11:
        raise ValueError("the August 20 template must contain exactly 11 slides")
    if presentation.slide_width != 9_144_000 or presentation.slide_height != 6_858_000:
        raise ValueError("the August 20 template must remain 10 x 7.5 inches")

    for slide in list(presentation.slides)[1:]:
        clear_slide_body(slide)

    builders = [
        build_title_slide,
        build_feedback_slide,
        build_progress_slide,
        build_label_chain_slide,
        build_snapshot_slide,
        build_questions_slide,
        build_decision_rules_slide,
        build_novelty_slide,
        build_reproducibility_slide,
        build_next_work_slide,
        build_appendix_slide,
    ]
    for slide, builder, notes in zip(
        presentation.slides, builders, SLIDE_NOTES, strict=True
    ):
        builder(slide)
        add_speaker_notes(slide, notes)

    properties = presentation.core_properties
    properties.author = "Krti Tallam"
    properties.last_modified_by = "Krti Tallam"
    properties.title = "Praxis Advisor Update - September 3, 2026"
    properties.subject = "Current Research Implementation and Evidence Controls"
    properties.comments = ""
    properties.keywords = ""
    properties.category = ""
    properties.content_status = ""
    properties.identifier = ""
    properties.language = "en-US"
    properties.version = ""

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    sanitize_inherited_package_metadata(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_deck())
