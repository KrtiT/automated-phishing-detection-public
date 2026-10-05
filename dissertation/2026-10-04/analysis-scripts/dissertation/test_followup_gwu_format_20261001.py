import importlib.util
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
SPEC = importlib.util.spec_from_file_location(
    "reading_builder", HERE / "build_followup_working_20261001.py"
)
BUILDER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILDER)


class GWUManuscriptFormatTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary = TemporaryDirectory()
        original_here, original_output = BUILDER.HERE, BUILDER.OUTPUT
        try:
            BUILDER.HERE = Path(cls.temporary.name)
            BUILDER.OUTPUT = BUILDER.HERE / "manuscript.docx"
            BUILDER.main()
            cls.document = Document(BUILDER.OUTPUT)
        finally:
            BUILDER.HERE, BUILDER.OUTPUT = original_here, original_output

    @classmethod
    def tearDownClass(cls):
        cls.temporary.cleanup()

    def test_template_page_geometry(self):
        for section in self.document.sections:
            self.assertEqual(section.page_width.inches, 8.5)
            self.assertEqual(section.page_height.inches, 11)
            self.assertEqual(section.left_margin.inches, 1.25)
            self.assertEqual(section.right_margin.inches, 1.25)
            self.assertEqual(section.top_margin.inches, 1)
            self.assertEqual(section.bottom_margin.inches, 1)

    def test_front_roman_body_arabic_numbering(self):
        self.assertEqual(len(self.document.sections), 2)
        front, body = self.document.sections
        self.assertEqual(front._sectPr.find(qn("w:pgNumType")).get(qn("w:fmt")), "lowerRoman")
        self.assertTrue(front.different_first_page_header_footer)
        self.assertEqual(body._sectPr.find(qn("w:pgNumType")).get(qn("w:start")), "1")
        self.assertFalse(body.different_first_page_header_footer)

    def test_body_font_spacing_and_chapter_breaks(self):
        chapters = [paragraph for paragraph in self.document.paragraphs if paragraph.text.startswith("Chapter ") and paragraph.style.name == "Heading 1"]
        self.assertEqual(len(chapters), 5)
        for chapter in chapters[1:]:
            self.assertTrue(chapter.paragraph_format.page_break_before)
        normal = self.document.styles["Normal"]
        self.assertEqual(normal.font.name, "Times New Roman")
        self.assertEqual(normal.font.size.pt, 12)
        paragraphs = self.document.paragraphs
        start = next(index for index, paragraph in enumerate(paragraphs) if paragraph is not None and paragraph.text == "Chapter 1—Introduction" and paragraph.style.name == "Heading 1")
        body = [paragraph for paragraph in paragraphs[start:] if paragraph.text.strip() and paragraph.style.name == "Normal"]
        self.assertTrue(body)
        for paragraph in body:
            if not paragraph.text.startswith("Table "):
                self.assertEqual(paragraph.paragraph_format.line_spacing, 2)

    def test_current_front_matter_preserves_personal_wording(self):
        text = "\n".join(paragraph.text for paragraph in self.document.paragraphs)
        for expected in ["The George Washington University", "Dedication", "Acknowledgements", "Abstract of Praxis", "Table of Contents", "List of Tables", "List of Symbols", "List of Acronyms"]:
            self.assertIn(expected, text)
        self.assertIn(BUILDER.ABSTRACT, text)
        for obsolete in ["Results-complete dissertation reading edition", "has passed the Final Examination", "final and approved form", "2,111,133 URL records", "Lorem ipsum"]:
            self.assertNotIn(obsolete, text)

    def test_paginated_navigation_and_caption_links(self):
        instructions = self.document._element.xpath(".//w:instrText")
        self.assertGreater(sum("PAGEREF" in (node.text or "") for node in instructions), 40)
        captions = [paragraph for paragraph in self.document.paragraphs if paragraph.style.name == "Caption"]
        self.assertEqual(len(captions), len(self.document.tables) + 4)
        for paragraph in captions:
            if paragraph.text.startswith("Table "):
                self.assertTrue(paragraph.paragraph_format.keep_with_next)
            self.assertEqual(paragraph.paragraph_format.first_line_indent, 0)

    def test_architecture_figure_and_roman_front_references(self):
        self.assertEqual(len(self.document.inline_shapes), 4)
        text = "\n".join(paragraph.text for paragraph in self.document.paragraphs)
        self.assertIn("List of Figures", text)
        anchors = self.document._element.xpath('.//w:hyperlink[@w:anchor="dedication"]')
        self.assertEqual(len(anchors), 1)

    def test_no_break_only_paragraphs_create_blank_pages(self):
        for paragraph in self.document.paragraphs:
            if not paragraph.text.strip():
                self.assertFalse(paragraph._element.xpath('.//w:br[@w:type="page"]'))


if __name__ == "__main__":
    unittest.main()
