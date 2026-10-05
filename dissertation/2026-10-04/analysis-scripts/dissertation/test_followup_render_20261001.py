import unittest
from pathlib import Path

import build_followup_working_20261001 as builder
import check_followup_render_20261001 as checker
import fitz
from docx import Document


class FollowupRenderTests(unittest.TestCase):
    def test_embedded_times_new_roman_is_accepted(self):
        with fitz.open() as document:
            page = document.new_page()
            page.insert_font(
                fontname="Manuscript",
                fontfile=str(Path("/System/Library/Fonts/Supplemental/Times New Roman.ttf")),
            )
            page.insert_text((90, 90), "Test manuscript", fontname="Manuscript")
            fonts = checker.check_fonts(document)
            self.assertTrue(fonts)
            self.assertTrue(all(font["embedded"] for font in fonts))

    def test_substituted_font_is_rejected(self):
        with fitz.open() as document:
            document.new_page().insert_text((90, 90), "Test manuscript", fontname="helv")
            with self.assertRaisesRegex(ValueError, "font"):
                checker.check_fonts(document)

    def test_matching_field_and_front_link_caches_are_accepted(self):
        document = Document()
        pages = {"heading_1": "1", "abstract": "vi"}
        for name, kind in [("heading_1", "body"), ("abstract", "front")]:
            builder.index_entry(document, {"text": name, "bookmark": name, "kind": kind, "level": 1}, pages)
        self.assertEqual(checker.check_cached_navigation(document, pages), 2)

    def test_stale_field_cache_is_rejected(self):
        document = Document()
        builder.field(document.add_paragraph(), " PAGEREF heading_1 \\h ", "9")
        with self.assertRaisesRegex(ValueError, "navigation"):
            checker.check_cached_navigation(document, {"heading_1": "1"})

    def test_missing_navigation_cache_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "navigation"):
            checker.check_cached_navigation(Document(), {"heading_1": "1"})

    def test_abstract_is_concise_without_omitting_pending_status(self):
        self.assertLessEqual(len(builder.ABSTRACT.split()), 215)
        self.assertIn("pending", builder.ABSTRACT)
        self.assertIn("H1–H3", builder.ABSTRACT)


if __name__ == "__main__":
    unittest.main()
