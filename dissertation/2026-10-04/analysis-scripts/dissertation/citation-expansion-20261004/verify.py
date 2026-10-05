"""Verify authorized text changes, retained results, citations and pagination."""

import csv
import json
import re
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

from docx import Document

import expand

sys.path.insert(0, str(expand.HERE.parent))
sys.path.insert(0, str(expand.CONTEXT / "gwu_advisor_update_source"))
import check_complete_package_20261002 as scientific
import check_followup_render_20261001 as render


def require(condition, message):
    if not condition:
        raise ValueError(message)


def prose(document):
    return [paragraph.text for paragraph in document.paragraphs
            if not paragraph._p.xpath(".//w:instrText|.//w:hyperlink[@w:anchor]")]


def rows(name):
    with (expand.CONTEXT / "deliverables/Tallam_Dissertation_Layout_Reviewed_2026-10-02/aggregate-data" / name).open() as stream:
        return list(csv.DictReader(stream))


def reference_audit(markdown):
    body, tail = markdown.split("## References\n\n", 1)
    references = tail.split("# Appendix A", 1)[0].strip().split("\n\n")
    require(len(references) == len(set(references)) == 58, "Expected 58 unique references")
    matched_spans = []
    report = []
    for reference in references:
        authors, year = re.match(r"^(.*?) \(((?:19|20)\d{2}[a-z]?)\)", reference).groups()
        first = authors.split(",", 1)[0]
        names = re.findall(r"(?:^|, & |, )([A-ZŽ][^,]+), [A-Z]", authors)
        stem = re.escape(first)
        if len(names) > 2:
            stem += r" et al\."
        elif len(names) == 2:
            stem += r" (?:and|&) " + re.escape(names[1])
        pattern = re.compile(r"\b" + stem + r"(?:['’]s)?\s*(?:\(\s*|,\s*)(?P<years>(?:19|20)\d{2}[a-z]?"
                             r"(?:,\s*(?:19|20)\d{2}[a-z]?)*)(?=[;).,\s])")
        matches = [match for match in pattern.finditer(body) if year in re.findall(r"(?:19|20)\d{2}[a-z]?", match["years"])]
        require(matches, "Uncited reference: " + reference)
        matched_spans.extend((match.start(), match.end()) for match in matches)
        report.append({"reference": reference, "citation_count": len(matches),
                       "citations": [{"line": body.count("\n", 0, match.start()) + 1, "text": match.group()} for match in matches]})
    residual = [token.group() for token in re.finditer(r"(?<!\d)(?:19|20)\d{2}[a-z]?(?!\d)", body)
                if not any(start <= token.start() and token.end() <= end for start, end in matched_spans)]
    prior = json.loads((expand.BASE / "citation-audit.json").read_text())
    require(residual == [entry["year"] for entry in prior["year_contexts_for_manual_review"]], "New unmatched year context")
    (expand.ROOT / "citation-audit.json").write_text(json.dumps({"references": report, "residual_year_tokens": residual,
        "residual_review": "Same 18 date/identifier/decimal tokens reviewed in the preserved consistency package; no new residuals."}, indent=2, ensure_ascii=False) + "\n")
    return len(references), sum(entry["citation_count"] for entry in report)


def main():
    baseline = Document(expand.SOURCE)
    current = Document(expand.OUTPUT)
    expected = prose(baseline)
    expected_markdown = expand.SOURCE.with_suffix(".md").read_text()
    for edit in expand.EDITS:
        matches = [index for index, text in enumerate(expected) if edit["old"] in text]
        require(len(matches) == 1, "Ambiguous expected edit")
        index = matches[0]
        expected[index] = expected[index].replace(edit["old"], edit["new"], 1)
        expected_markdown = expected_markdown.replace(edit["old"], edit["new"], 1)
    for addition in expand.ADDITIONS:
        anchor = addition["before"].removeprefix("## ")
        matches = [index for index, text in enumerate(expected) if text.startswith(anchor)]
        require(len(matches) == 1, "Ambiguous expected insertion")
        expected.insert(matches[0], addition["text"])
        expected_markdown = expected_markdown.replace(addition["before"], addition["text"] + "\n\n" + addition["before"], 1)
    for reference in expand.REFERENCES:
        text = reference["prefix"] + reference["italic"] + reference["suffix"]
        index = next(index for index, item in enumerate(expected) if item.startswith(reference["before"]))
        expected.insert(index, text)
        expected_markdown = expected_markdown.replace("\n\n" + reference["before"], "\n\n" + text + "\n\n" + reference["before"], 1)
        paragraph = expand.target(current, text)
        require(paragraph.paragraph_format.left_indent.inches == .5, "Reference left indent")
        require(paragraph.paragraph_format.first_line_indent.inches == -.5, "Reference hanging indent")
        require([run.text for run in paragraph.runs if run.italic] == [reference["italic"]], "Reference emphasis")
    markdown = expand.OUTPUT.with_suffix(".md").read_text()
    require(markdown == expected_markdown and prose(current) == expected, "Unlogged text change")
    require([paragraph._p.xml for paragraph in current.paragraphs[:24]] ==
            [paragraph._p.xml for paragraph in baseline.paragraphs[:24]], "Front matter or credentials changed")
    require(len(current.tables) == 27 and len(current.inline_shapes) == 4, "Table/figure inventory")
    require([table._tbl.xml for table in current.tables] == [table._tbl.xml for table in baseline.tables], "Table changed")
    with ZipFile(expand.SOURCE) as original, ZipFile(expand.OUTPUT) as revised:
        for name in original.namelist():
            if name.startswith("word/media/"):
                require(original.read(name) == revised.read(name), "Figure bytes changed")
    prior_markdown = expand.SOURCE.with_suffix(".md").read_text()
    for start, end in [("# Chapter 1—", "# Chapter 2—"), ("# Chapter 4—", "# Chapter 5—"), ("# Appendix A—", None)]:
        old = prior_markdown[prior_markdown.index(start):]
        new = markdown[markdown.index(start):]
        if end:
            old, new = old.split(end, 1)[0], new.split(end, 1)[0]
        require(old == new, "Changed protected scientific section: " + start)
    require(not current._element.xpath(".//w:ins|.//w:del"), "Tracked changes remain")
    require(not re.search(r"\b(?:TODO|TBD|FIXME)\b|Error! Reference", markdown), "Unresolved placeholder")
    refs, links = reference_audit(markdown)
    gates = rows("primary-gates.csv")
    displayed = [[[cell.text for cell in row.cells] for row in table.rows] for table in current.tables]
    require(len(gates) == len(displayed[16][1:]) == 22, "Primary check inventory")
    for saved, actual in zip(gates, displayed[16][1:]):
        unit = " ms" if saved["name"] == "http_pooled_p95_ms" else " pp" if "_minus_" in saved["name"] else "%"
        scale = 1 if unit == " ms" else 100
        operand = f"{float(saved['estimate']) * scale:.4f}{unit}"
        threshold = f"{saved['operator']} {float(saved['threshold']) * scale:.4f}{unit}"
        if unit == " ms":
            threshold = f"{saved['operator']} {float(saved['threshold']):g}{unit}"
        require(actual == [saved["hypothesis"], saved["name"].replace("_", " "), operand, threshold, saved["status"]], "Primary result mismatch")
    groups = rows("operational-groups.csv")
    require(len(groups) == 25 and len(rows("operational-runs.csv")) == 125, "Operational inventory")
    for saved, actual in zip(groups, displayed[8][1:]):
        values = [f"{float(saved[field]):.3f}" for field in ["p50_ms", "p95_ms", "p99_ms"]]
        values += [f"{saved['request_errors']}/{saved['request_count']}", f"{100 * float(saved['physical_invocation_fraction']):.4f}%"]
        require(actual[1:] == values, "Operational result mismatch")
    numeric = scientific.check_numeric_tables(current)
    preserved = json.loads((expand.HERE / "preservation-before.json").read_text())
    for name, expected_hash in preserved.items():
        require(expand.digest(Path(name)) == expected_hash, "Protected input changed: " + name)
    for path in (expand.ROOT / "advisor").iterdir():
        require(expand.digest(path) == expand.digest(expand.BASE / "advisor" / path.name), "Advisor deck changed")
    render.HERE = render.REVIEW = expand.MANUSCRIPT
    render.PDF = expand.OUTPUT.with_suffix(".pdf")
    render.main()
    report = {
        "status": "citation_expansion_verified", "checked_at": datetime.now(timezone.utc).isoformat(),
        "references": refs, "new_primary_references": 7, "exact_reference_citation_links": links,
        "authorized_span_edits": 8, "added_literature_paragraphs": 2,
        "primary_checks": 22, "primary_decisions": dict(Counter(item["status"] for item in gates)),
        "operational_groups": 25, "operational_cells": 125, "unchanged_tables": 27, "unchanged_figures": 4,
        "other_numeric_checks": numeric, "chapters_1_4_and_appendices_unchanged": True,
        "front_matter_and_credentials_unchanged": True, "advisor_deck_unchanged": True,
        "protected_input_hashes": preserved,
        "outputs": {path.name: expand.digest(path) for path in [expand.OUTPUT, expand.OUTPUT.with_suffix(".pdf"), expand.OUTPUT.with_suffix(".md")]},
        "limits": "Not AIR/Turnitin clearance, confirmation of committee roles, an institutional page-length waiver or a full-source audit of all existing references.",
    }
    (expand.ROOT / "verification.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
