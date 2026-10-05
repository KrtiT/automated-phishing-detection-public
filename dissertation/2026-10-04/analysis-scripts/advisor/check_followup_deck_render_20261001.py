"""Check the 23-slide native export and retain slide views and speaker notes."""

import hashlib
import json
import re
import unicodedata

import fitz
from PIL import Image, ImageDraw
from pptx import Presentation

from build_followup_deck_20261001 import OUTPUT

EXPECTED_SLIDES = 23


def normalize(value):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", value).replace("\u00ad", ""))


def text_items(slide):
    for shape in slide.shapes:
        if shape.has_text_frame:
            yield from (paragraph.text for paragraph in shape.text_frame.paragraphs if paragraph.text.strip())
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield from (paragraph.text for paragraph in cell.text_frame.paragraphs if paragraph.text.strip())


def check():
    presentation = Presentation(OUTPUT)
    missing, outside, checked = [], [], 0
    with fitz.open(OUTPUT.with_suffix(".pdf")) as rendered:
        if len(presentation.slides) != EXPECTED_SLIDES or len(rendered) != EXPECTED_SLIDES:
            raise ValueError(f"Expected all {EXPECTED_SLIDES} slides")
        for index, (slide, page) in enumerate(zip(presentation.slides, rendered), 1):
            pdf_text = normalize(page.get_text())
            if not pdf_text:
                raise ValueError(f"Empty rendered slide {index}")
            for text in text_items(slide):
                checked += 1
                if normalize(text) not in pdf_text:
                    missing.append({"slide": index, "text": text})
            for word in page.get_text("words"):
                if word[0] < 0 or word[1] < 0 or word[2] > page.rect.width or word[3] > page.rect.height:
                    outside.append({"slide": index, "word": word[:5]})
    return {"slides": EXPECTED_SLIDES, "checked_text_items": checked, "missing_text": missing,
            "out_of_bounds_words": outside,
            "pptx_sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
            "pdf_sha256": hashlib.sha256(OUTPUT.with_suffix(".pdf").read_bytes()).hexdigest(),
            "render_engine": "LibreOffice 26.8.0.3 local installation",
            "status": "mechanical_checks_passed" if not missing and not outside else "review_required"}


def main():
    report = check()
    directory = OUTPUT.parent
    presentation = Presentation(OUTPUT)
    with fitz.open(OUTPUT.with_suffix(".pdf")) as rendered:
        for index, page in enumerate(rendered, 1):
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(directory / f"slide-{index:02}.png")
        for start in range(0, len(rendered), 6):
            sheet = Image.new("RGB", (1500, 1410), "#dddddd")
            drawing = ImageDraw.Draw(sheet)
            for offset, index in enumerate(range(start, min(start + 6, len(rendered)))):
                picture = Image.open(directory / f"slide-{index + 1:02}.png")
                picture.thumbnail((720, 425))
                left, top = 15 + 750 * (offset % 2), 30 + 470 * (offset // 2)
                sheet.paste(picture, (left, top))
                drawing.text((left, top - 20), f"Slide {index + 1}", fill="black")
            sheet.save(directory / f"contact-{start // 6 + 1:02}.png")
    notes = [f"SLIDE {index}\n{slide.notes_slide.notes_text_frame.text.strip()}"
             for index, slide in enumerate(presentation.slides, 1)]
    notes_path = OUTPUT.with_name(OUTPUT.stem + "_Speaker_Notes.txt")
    notes_path.write_text("\n\n".join(notes) + "\n")
    report["speaker_notes_sha256"] = hashlib.sha256(notes_path.read_bytes()).hexdigest()
    report["visual_review"] = "Views generated; inspection is separate."
    (directory / "render-check.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))
    if report["status"] != "mechanical_checks_passed":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
