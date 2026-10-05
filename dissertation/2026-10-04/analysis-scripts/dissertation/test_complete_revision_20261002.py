import importlib
import unittest


class CompleteRevisionTests(unittest.TestCase):
    def setUp(self):
        self.revision = importlib.import_module("build_complete_revision_20261002")

    def test_preserves_original_questions_and_results(self):
        body = self.revision.compose_body()
        original = self.revision.ORIGINAL.read_text()
        for start,end in [("## 1.4 Research Questions and Hypotheses","## 1.5 Contribution Boundary"),
                          ("## 4.1 Completed Evaluation and Populations","# Chapter 5—Discussion and Conclusions")]:
            self.assertIn(original.split(start,1)[1].split(end,1)[0].strip(),body)
        self.assertEqual(body.count("# Chapter "),5)

    def test_all_service_results_and_chronology_are_present(self):
        body = self.revision.compose_body()
        for text in ["800,000","80,000","99,999","99,200","799","0.203249","71.9990",
                     "184.66","414","45 completed arms","not supported","Table 4.18.","Table 4.19.","Table 4.20.","Table B.1."]:
            self.assertIn(text,body)
        self.assertNotIn("S remains unmeasured",body)
        self.assertNotIn("service comparison remains pending",body)
        self.assertNotIn("no controlled measurements yet",body)
        self.assertIn("independent arithmetic",body)
        self.assertIn("one shared-client",body)

    def test_abstract_states_real_improvement_and_full_decision(self):
        abstract = self.revision.abstract()
        self.assertLessEqual(len(abstract.split()),215)
        for text in ["79.68%","72.00 ms","H1–H3","S","sequence"]:
            self.assertIn(text,abstract)
        self.assertNotIn("pending",abstract)

    def test_all_eighty_arms_have_appendix_rows(self):
        body = self.revision.compose_body()
        rows = [line for line in body.splitlines() if line.startswith("| no_model-c") or line.startswith("| structural_detector-c")]
        self.assertEqual(len(rows),80)

    def test_representation_conformance_remains_separate_from_efficacy(self):
        self.assertIn("not an additional efficacy endpoint.", self.revision.compose_body())

    def test_no_pending_service_claims_in_final_speaker_notes(self):
        notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in self.revision.compose_deck().slides)
        for stale in ("S is not yet measured", "service comparison pending", "S remains unmeasured"):
            self.assertNotIn(stale, notes)

    def test_deck_preserves_original_promises_and_twenty_two_checks(self):
        from pptx import Presentation
        import build_followup_deck_20261001 as previous
        current = self.revision.compose_deck()
        original = Presentation(previous.ORIGINAL)
        def texts(slide):
            return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame) + "\n".join(cell.text for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells)
        self.assertEqual(len(current.slides),25)
        self.assertEqual(texts(current.slides[1]),texts(original.slides[1]))
        for original_index, current_index in zip(range(11,17),range(19,25)):
            self.assertEqual(texts(current.slides[current_index]),texts(original.slides[original_index]))
        all_text = "\n".join(texts(slide) for slide in current.slides)
        self.assertNotIn("not yet measured",all_text)
        self.assertIn("79.68%",all_text)


if __name__ == "__main__":
    unittest.main()
