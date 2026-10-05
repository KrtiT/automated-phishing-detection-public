"""Integrate the verified, complete bounded comparison without altering prior copies."""

import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Inches

import build_followup_working_20261001 as manuscript
import build_followup_deck_20261001 as deck
from integrate_followup_manuscript_20261001 import replace_once

HERE = Path(__file__).resolve().parent
ROOT = HERE / "followup-20261001"
ORIGINAL = HERE / "Tallam_Krti_Praxis_v3_Body_Results_2026-10-01.md"
PREVIOUS = ROOT / "manuscript-work/Tallam_Krti_Praxis_Engineering_Working_2026-10-01.md"
FINAL = ROOT / "complete-revision-20261002"
BODY = FINAL / "manuscript/Tallam_Krti_Praxis_GWU_Complete_2026-10-02.md"
DOCX = BODY.with_suffix(".docx")
PPTX = FINAL / "advisor/Tallam_Praxis_Advisor_Complete_2026-10-02.pptx"


def evidence():
    path = ROOT / "verified-service-v2/verification.json"
    report = json.loads(path.read_bytes())
    if report["status"] != "verified" or report["arms"] != 80 or report["measured_requests"] != 800000:
        raise ValueError("Complete verified service evidence is required")
    for name, expected in report["aggregate_sha256"].items():
        if hashlib.sha256((path.parent / name).read_bytes()).hexdigest() != expected:
            raise ValueError("Verified service aggregate changed")
    return report


def rows(name):
    with (ROOT / "verified-service-v2" / name).open() as stream:
        return list(csv.DictReader(stream))


def table(headers, records):
    return "| " + " | ".join(headers) + " |\n|" + "---|" * len(headers) + "\n" + "\n".join("| " + " | ".join(str(value) for value in row) + " |" for row in records) + "\n\n"


def abstract():
    evidence()
    return """This study develops a raw-URL inference gateway and evaluates representation value, Gaussian-mixture-guided escalation under source/domain shift, and inline detection–service tradeoffs. The initial study includes 34,593 internal records, 8,701 external records, 125 HTTP cells and 22 primary checks. Nine component checks pass; the original H1–H3 conjunctions are not supported. Structural features improve internal recall by 64.07 percentage points over length alone, while external specificity and useful escalation remain unresolved. Diagnosis motivates two bounded modifications. A transport-neutral structural model achieves exact feature and score invariance across 8,622 prescribed HTTP/HTTPS companion pairs. At frozen thresholds it yields 251 fewer false positives and 184 fewer detected positives; detection requirement D is not met. An 80-arm service comparison evaluates worker-owned persistent HTTP clients against a shared pool. At concurrency 64, the structural workload's median paired p95 falls by 79.68% (ratio 0.20325; 97.5% interval [0.18790, 0.21351]); worker p95 is 72.00 ms with zero errors in 100,000 requests. Prediction fields agree on all 99,999 comparable pairs. Admission-sequence differences and one shared-client timeout prevent strict S agreement. The contribution is a connected artifact and measured engineering process that separates representation stability, detection risk, routing utility and transport cost. Retrospective source labels, sampled host conditions and explicitly retained recovery constrain generalization."""


def service_results():
    report = evidence()
    primary = report["primary"]
    text = """## 4.9 Controlled Service Comparison: Completed Evaluation

### 4.9.1 Execution, Recovery and Complete Accounting

The complete authorized recovery schedule ran on October 2, 2026, from 07:49:30Z to 08:11:48Z. All 80 arms completed: two workloads, two concurrency levels, ten pairs and two client topologies. This adds 800,000 measured requests and 80,000 warmup requests, without pooling the earlier interrupted attempt. All 80 service processes exited normally without forced termination. The preflight retained 37 clean samples spanning 184.66 seconds; 414 schedule observations reported AC power, normal available thermal/performance status, owned sleep inhibition and no detected competing workload. These are sampled observations plus the operator's host reservation, not proof about every instant.

The first service attempt stopped on an OS-reported AC loss at 07:38:54Z after 45 completed arms and part of the next arm. Its 450,000 completed measured requests had no recorded request errors. No primary structural concurrency-64 arm had begun. A disclosed amendment authorized one new complete schedule after stable AC. All 372 first-attempt files, including its interruption and partial progress, remain hash-identical and excluded from the recovery result. This change followed partial nonprimary observations and is not presented as a pre-experiment registration.

Independent arithmetic recomputation from saved JSON verifies all 80 arm summaries, 40 paired comparisons, eight pooled groups and five S requirements against the frozen reducer. It performs no predictions or fitting and is not independent-investigator replication. The new schedule records one shared-client timeout and zero warmup errors. The original study's 125 cells, 25 groups, 22 primary checks and 1,901 recorded errors remain separate and unchanged.

### 4.9.2 Paired Tail-Latency Effect

Table 4.18 gives every primary pair. Every worker p95 is lower than its paired shared-client value. The median worker/shared ratio is 0.203249, corresponding to a 79.68% median paired reduction. The two-sided 97.5% run-bootstrap interval is [0.187903, 0.213508], from the fixed 10,000 resamples; no pair or replicate is undefined. The pooled worker success p95 is 71.9990 ms, below the 200-ms requirement. The pooled shared success p95 is 358.3663 ms. A ratio of pooled quantiles is not substituted for the prespecified median of paired ratios.

Table 4.18. All ten structural-detector concurrency-64 pairs; success-only p95 in milliseconds.

"""
    text += table(["Pair","Shared p95","Worker p95","Ratio","Exact / requested"],
                  [[row["pair"],f'{float(row["shared_p95_ms"]):.4f}',f'{float(row["worker_p95_ms"]):.4f}',f'{float(row["ratio"]):.6f}',f'{row["exact_except_request_id"]} / 10,000'] for row in rows("primary-pairs.csv")])
    text += """### 4.9.3 Controls, Throughput, Errors and Response Semantics

Table 4.19 reports all required control/concurrency combinations, not only the primary comparison. Each row pools ten arms and 100,000 measured attempts. Throughput is attempted requests divided by the sum of measured client-phase durations; it excludes separately retained warmup and drain intervals. With zero errors, successful-response throughput equals attempted throughput. The shared structural c64 row instead has 99,999 successes and one error. Successful, failed and all-request percentiles and physical counters are preserved in the aggregate CSVs; Appendix B gives every arm.

Table 4.19. Complete pooled service groups; success p50/p95/p99 in milliseconds and attempted throughput in requests per second.

"""
    text += table(["Workload / c / client","p50","p95","p99","Errors","Requests/s"],
                  [[f'{"Control" if row["workload"] == "no_model" else "Structural"} / {row["concurrency"]} / {row["client"]}',*[f'{float(row[key]):.2f}' for key in ("success_p50_ms","success_p95_ms","success_p99_ms")],row["request_errors"],f'{float(row["client_attempts_per_second"]):.2f}'] for row in rows("group-metrics.csv")])
    text += """At concurrency 1, the two topologies have nearly identical pooled p95: 1.10 versus 1.11 ms for the control and 1.42 versus 1.43 ms for the structural workload. At concurrency 64, control p95 falls from 345.95 to 67.21 ms and structural p95 from 358.37 to 72.00 ms. Structural attempted throughput rises from 454.14 to 1,526.71 requests/s; successful shared throughput is 454.14 requests/s when rounded to two decimals. The control pattern supports a client/transport explanation for this intervention's benefit rather than a faster classifier. It does not isolate individual connection-pool, scheduling or queueing mechanisms, and it does not estimate a transformer speedup.

The only request error occurs in structural c64 pair 2's shared arm. Its timeout latency is 2,019.144208 ms; that is the sole observation for all three failure-only quantiles. A two-second deadline can produce a terminal elapsed time slightly above two seconds when cancellation and bookkeeping complete. This is retained as a timeout, not relabeled a success. That arm admits and completes 9,999 measured requests; the timed-out request is not counted as admitted. Across the full schedule there are 799,999 measured admissions/completions, zero server failed-request increments and zero transformer attempts or successful transformer scores. A client timeout and a server failure are different counters.

The 100,000 primary worker requests all succeed. Of the paired shared/worker requests, 99,999 are comparable and all agree exactly on action, probability and stage-2 invocation after excluding request IDs and admission sequence. The timeout leaves one pair noncomparable. Strict agreement excludes request IDs only: 799/100,000 primary pairs satisfy it. The remaining 99,200 comparable pairs differ in admission sequence, not prediction fields. Across all 40 pairs, prediction agreement is 399,999/399,999 comparable pairs, with the same one noncomparable error. This does not justify saying that all 400,000 requested pairs have observed agreement.

### 4.9.4 Requirement S and Engineering Interpretation

Table 4.20 separates demonstrated performance from the unchanged response contract. Four of five S requirements pass. S as a conjunction is not supported because full response equivalence includes admission sequence. No threshold or agreement field is removed after seeing this result.

Table 4.20. Prespecified S requirements and verified decisions.

"""
    values = [
        ["Median paired ratio ≤0.8",f'{primary["median_ratio"]:.6f}',"Pass"],
        ["97.5% ratio upper bound <1",f'{primary["ratio_interval_97_5"][1]:.6f}',"Pass"],
        ["Pooled worker success p95 ≤200 ms",f'{primary["worker_success_latency"]["p95_ms"]:.4f} ms',"Pass"],
        ["Worker errors <0.1%", "0 / 100,000 (0%)","Pass"],
        ["Exact paired responses except IDs","799 / 100,000","Not met"],
    ]
    text += table(["Requirement","Observed result","Decision"],values)
    text += """The modification nevertheless demonstrates a concrete service improvement: the structural HTTP workload meets the tail-latency and reliability requirements with worker-owned clients, while preserving every comparable prediction. This is a measured transport improvement, not a deployment qualification for the detector and not retrospective support for original H3. Applications that require equivalent admission ordering would need an explicitly designed ordering contract and a new evaluation. Applications that need request-correlated prediction equivalence may consider these results informative, but cannot retroactively substitute that narrower contract for S.

"""
    return text


def compose_body():
    evidence()
    body = PREVIOUS.read_text()
    start = body.index("## 4.9 Controlled Service Comparison: Working-Copy Status")
    end = body.index("# Chapter 5—Discussion and Conclusions",start)
    body = body[:start] + service_results() + body[end:]
    replacements = {
        "The controlled service comparison remains pending in this working copy.":"The completed service comparison demonstrates a substantial tail-latency improvement while separately testing strict response equivalence.",
        "A no-model service control and paired client-topology design isolate an additional transport hypothesis, with final measurements still pending in this working copy.":"A no-model control and paired client-topology comparison establish a 79.68% median paired p95 reduction in the structural concurrency-64 workload. Worker-owned clients meet its latency and error requirements; admission-order equivalence remains unmet.",
        "candidate observed FPR at most 1%, and exact candidate scheme invariance.":"candidate observed FPR at most 1%. Exact candidate scheme invariance is a separate required representation-conformance check, not an additional efficacy endpoint.",
        "S remains unmeasured in this working copy and supports no performance or agreement claim.":"S is measured on one reserved Mac with synthetic strings, a singleton structural scorer and ten operational pairs. It does not establish production arrival-rate capacity, transformer performance, field prevalence or cross-hardware generalization. The interrupted schedule and disclosed full-schedule recovery are retained separately. Sampled condition checks cannot exclude every unobserved disturbance. Response-order differences are not prediction differences, and the one timeout prevents complete paired comparability.",
        "The controlled service comparison remains pending; this working copy makes no final S claim.":"The service modification yields a 79.68% median paired p95 reduction, a 72.00-ms pooled worker p95 and zero primary worker errors. All 99,999 comparable prediction pairs agree. Strict S remains not supported because admission-sequence equivalence is not achieved and one shared timeout is noncomparable; four of its five prespecified requirements pass.",
        "The final extension package must add the controlled service evidence and complete its cross-document verification before it is labeled complete.":"The complete initial and follow-up evidence is retained with the final manuscript, editable advisor deck, aggregate analysis and verification records. Research completion does not assert institutional acceptance or general deployment readiness.",
        "| Follow-up S: client-topology service comparison | Fixed 80-arm comparison is implemented; no controlled measurements yet | Unmeasured in this working copy; no performance claim |":"| Follow-up S: client-topology service comparison | Four of five requirements pass: median ratio 0.203249; interval upper 0.213508; worker p95 71.9990 ms; zero worker errors | Not supported as a conjunction; strict agreement is 799/100,000. All 99,999 comparable predictions agree |",
    }
    for old,new in replacements.items():
        body = replace_once(body,old,new)
    methods = """A dated October 2 recovery amendment superseded only the initial no-additional-schedule rule after an observed power interruption. The first 45 complete arms and interrupted 46th arm were preserved; no primary structural c64 arm had begun. Following the operator's explicit approval, a separate hash-bound launcher reused the unchanged frozen source, all 80 arms, thresholds, deadlines and strict response semantics. It required at least 180 seconds of sampled stable AC and retained the original continuous guard loop during one new complete schedule. No partial arm or complete arm from the first attempt contributes to the primary S result. The amendment was made after partial nonprimary observations, not backdated, and asserted no advisor or institutional approval.

"""
    body = replace_once(body,"# Chapter 4—Results",methods+"# Chapter 4—Results")
    rq3 = """The follow-up makes the operational answer more precise. With the scorer unchanged, worker-owned clients reduce the structural c64 median paired p95 by 79.68%, place pooled worker p95 at 72.00 ms and complete 100,000 primary worker requests without errors. The no-model control improves similarly while c1 changes little, locating a demonstrated benefit in the client/transport path rather than classifier acceleration. Every comparable prediction remains identical. Strict admission ordering and the original detector's external-risk requirements remain separate unmet conditions. Thus an inline latency repair is demonstrated for the specified synthetic workload, but not an end-to-end low-FPR deployment or revised H3 decision.

"""
    body = replace_once(body,"## 5.2 End-to-End Argument and Contribution",rq3+"## 5.2 End-to-End Argument and Contribution")
    practical = """The service iteration supplies a complementary positive result. When the initial latency budget was not met, the investigation separated no-model HTTP work from structural scoring, changed only client connection ownership and measured the full paired schedule. The measured improvement is large and consistent across all ten primary pairs, not a best-run selection. At the same time, observing prediction agreement alongside admission-order disagreement exposes two different meanings of behavioral preservation. The resulting engineering guidance is to measure the client path and specify whether ordering is a required externally visible property; neither model-call economy nor prediction equality alone proves the entire service contract.

"""
    body = replace_once(body,"## 5.3 Implications for Practice",practical+"## 5.3 Implications for Practice")
    body += """

# Appendix B—Complete Follow-up Service Evidence

The following 80 rows preserve schedule order. Each arm has 1,000 warmups and 10,000 measured attempts. IDs identify workload, concurrency, pair and client; structural_detector is the unchanged structural scorer, and no_model is the response control. Successful-response p95 is reported here; the machine-readable arm file also provides successful, failed and all-request p50/p95/p99, measured/drain durations, error categories and physical counts. All eight pooled groups and all 40 pair-agreement rows accompany the package. Empty failure quantiles mean zero failures, not zero-millisecond failures.

Table B.1. Complete 80-arm schedule; success-only p95 in milliseconds, attempted throughput in requests/s and terminal error counts.

"""
    body += table(["Arm ID","Success p95","Requests/s","Errors"],
                  [[f'{row["workload"]}-c{row["concurrency"]}-pair{int(row["pair"]):02d}-{row["client"]}',f'{float(row["success_latency_p95_ms"]):.4f}',f'{float(row["client_attempts_per_second"]):.2f}',row["request_errors"]] for row in rows("arm-metrics.csv")])
    body += """The service source is the unchanged ef8ba5f0b357cf3dd60c4d663e6297d13334460c checkout. The recovery authorization binds the separate launcher, tests, amendment, original manifest and interrupted-attempt preservation receipt. The independent arithmetic verifier uses separately implemented sorted linear quantiles, medians, paired-response comparisons, physical-count reconciliation and the fixed PCG64 bootstrap. Its verification record binds the retained sources and five aggregate CSVs. It is a second calculation, not an external replication. All failure observations, including the original power interruption, shared-client timeout and unmet strict agreement, are retained. No fitting, prediction or measurement is performed during document synthesis.
"""
    return body


def compose_deck():
    report = evidence()
    presentation = Presentation(deck.OUTPUT)
    base = deck.base
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.text = run.text.replace("October 1, 2026","October 2, 2026").replace("Engineering working revision","Complete research revision")
        notes = slide.notes_slide.notes_text_frame
        notes.text = notes.text.replace("S is not yet measured; no diagnostic timing is final evidence.","S is now complete under the disclosed October 2 environmental-recovery amendment; original hypotheses and all thresholds remain unchanged.")
    for index in (14,15,16):
        base.LAYOUT.clear_slide_body(presentation.slides[index])
    slide = presentation.slides[14]
    base.panels(slide,"Measured service improvement; strict contract distinguished",
                "Primary structural workload, c64",["Median paired p95 ratio: 0.20325 (79.68% reduction).","97.5% interval: [0.18790, 0.21351].","Worker p95: 72.00 ms; zero errors / 100,000."],
                "Prediction versus response order",["99,999 / 99,999 comparable predictions agree.","One shared timeout leaves one pair noncomparable.","Strict agreement: 799 / 100,000; ordering differs."],
                "Four of five S requirements pass. S is not supported as a conjunction; original H3 is unchanged.")
    base.LAYOUT.add_speaker_notes(slide,["Complete 80-arm schedule; 800,000 measured and 80,000 warmup requests. One total measured error: structural c64 pair 2 shared timeout, 2,019.144208 ms. Zero warmup errors. Among 99,999 comparable primary pairs, 99,200 differ only in admission sequence. Exact response agreement excludes only request IDs. Prediction agreement additionally excludes sequence and is separately descriptive. All ten primary p95 ratios are lower than one. The 97.5% interval uses 10,000 resamples and seed 20261002. A measured service benefit does not cure external detection specificity or change original H3."])
    base.panels(presentation.slides[15],"Technical contribution: measured changes, explicit contracts",
                "What the artifact demonstrates",["Structural representation: +64.07 pp internal recall.","Scheme-neutral inputs: exact invariance on 8,622 pairs.","Worker-owned clients: 79.68% paired p95 reduction."],
                "What the evidence distinguishes",["Representation stability versus external low-FPR detection.","Shift sensitivity versus useful escalation after the alarm.","Prediction equivalence versus admission-order equivalence."],
                "Reusable artifact and evaluation contracts; no invention claim for normalization, regression or connection reuse.")
    base.LAYOUT.add_speaker_notes(presentation.slides[15],["The narrative is design → evaluation → diagnosis → modification → measured comparison. Original H1/H2/H3 remain unsupported; nine component checks pass. D remains unsupported, with exact invariance achieved. S passes four performance/reliability components but not strict response agreement. These findings support concrete, qualified engineering decisions rather than a claim that every joint requirement was met."])
    base.panels(presentation.slides[16],"Completed research and dissertation evidence",
                "Original commitments delivered",["All 125 cells, 25 operational groups and 22 primary checks.","Direct RQ1–RQ3 answers; explicit H1–H3 decisions.","Secondary analyses, raw-count tables and audit limits."],
                "Iteration and final artifacts",["D: all 8,622 rows and companions; no retuning.","S: all 80 arms; independent saved-evidence recomputation.","GWU manuscript, editable deck, aggregates and provenance."],
                "All adverse results and interruptions retained. Research completion is not institutional acceptance.")
    base.LAYOUT.add_speaker_notes(presentation.slides[16],["The original October 1 package, author manuscript, frozen source and interrupted evidence are preserved. On October 2 the operator explicitly authorized one full recovery schedule after transient AC loss. The first 45 arms never enter the primary S result. The new run retained 37 preflight samples over 184.66 seconds and 414 schedule samples, with zero recorded guard violations. No advisor or institutional approval is claimed. Deliverables retain the original credential wording and updated abstract/navigation."])
    base.LAYOUT.add_speaker_notes(presentation.slides[0],["The completed research combines the original evaluation and separately specified engineering iteration. Lead with the achieved representation and service properties, then state their scope and remaining joint-criterion failures. Original H1–H3, detection D and strict service S remain unsupported; there are nine passing original components and four passing S requirements. The package is research-complete, not a university submission certification."])
    added = [presentation.slides.add_slide(presentation.slides[1].slide_layout) for _ in range(2)]
    for slide in added:
        base.LAYOUT.clear_slide_body(slide)
    base.LAYOUT.set_title(added[0],"All ten primary pairs: unchanged structural scorer, c64")
    base.grid(added[0],["Pair","Shared p95 ms","Worker p95 ms","Ratio"],
              [[row["pair"],f'{float(row["shared_p95_ms"]):.2f}',f'{float(row["worker_p95_ms"]):.2f}',f'{float(row["ratio"]):.5f}'] for row in rows("primary-pairs.csv")],widths=[.13,.3,.3,.27],size=12,top=1.2,height=3.8)
    base.LAYOUT.add_note(added[0],"Every pair retained. Success-only p95; one shared timeout in pair 2 is retained separately. No best-run selection.",y=5.2,font_size=11)
    base.LAYOUT.add_speaker_notes(added[0],["Source: verified-service-v2/primary-pairs.csv. Each arm has 10,000 measured attempts after 1,000 warmups. Alternating client order, fresh service per arm. Bootstrap unit is the pair, not individual requests. Worker p95 ranges from 67.27 to 74.72 ms. Ratio estimates are conditional on this machine and synthetic workload."])
    base.LAYOUT.set_title(added[1],"Controls locate the measured benefit in the client path")
    base.grid(added[1],["Workload","c","Client","p95 ms","Errors","Req/s"],
              [["Control" if row["workload"] == "no_model" else "Structural",row["concurrency"],row["client"],f'{float(row["success_p95_ms"]):.2f}',row["request_errors"],f'{float(row["client_attempts_per_second"]):.1f}'] for row in rows("group-metrics.csv")],widths=[.21,.07,.16,.18,.14,.24],size=12,top=1.2,height=3.7)
    base.LAYOUT.add_note(added[1],"Ten arms and 100,000 attempts per row. Similar c1 latency; both c64 workloads improve. No transformer speedup claim.",y=5.2,font_size=11)
    base.LAYOUT.add_speaker_notes(added[1],["Source: verified-service-v2/group-metrics.csv. Attempted throughput divides total attempts by summed measured phase duration, excluding warmup and drains. Failure-only latency is undefined in seven groups because they have zero errors. The sole shared structural timeout is 2,019.144208 ms. Full success/failure/all-request quantiles and physical counters are exported. Control results do not isolate individual pool, scheduler or queueing mechanisms."])
    identifiers = list(presentation.slides._sldIdLst)
    for identifier in identifiers:
        presentation.slides._sldIdLst.remove(identifier)
    for identifier in identifiers[:15]+identifiers[23:]+identifiers[15:23]:
        presentation.slides._sldIdLst.append(identifier)
    for slide in list(presentation.slides)[1:]:
        title = slide.shapes.title
        title.left,title.top,title.width,title.height = Inches(.48),Inches(.22),Inches(9.02),Inches(.85)
    presentation.core_properties.title = "Praxis Complete Research Revision — October 2, 2026"
    presentation.core_properties.subject = "Original RQ1–RQ3 and completed D/S engineering comparisons"
    presentation.core_properties.modified = datetime.now(timezone.utc)
    if report["primary_prediction_agreement"] != 99999 or report["primary"]["exact_paired_agreement"] != 799:
        raise ValueError("Narrative no longer matches service evidence")
    return presentation


def main():
    BODY.parent.mkdir(parents=True,exist_ok=True)
    PPTX.parent.mkdir(parents=True,exist_ok=True)
    BODY.write_text(compose_body())
    manuscript.HERE,manuscript.BODY,manuscript.OUTPUT,manuscript.ABSTRACT = BODY.parent,BODY,DOCX,abstract()
    manuscript.main()
    document = Document(DOCX)
    for paragraph in document.paragraphs:
        if paragraph.text == "October 1, 2026":
            paragraph.runs[0].text = "October 2, 2026"
    document.core_properties.subject = "GWU D.Eng. Praxis — complete research and engineering comparisons"
    document.save(DOCX)
    metadata_path = BODY.parent/"results-pdf-metadata-20261001.json"
    metadata = json.loads(metadata_path.read_text())
    metadata["date"] = "October 2, 2026"
    metadata_path.write_text(json.dumps(metadata,indent=2)+"\n")
    compose_deck().save(PPTX)
    print(DOCX)
    print(PPTX)


if __name__ == "__main__":
    main()
