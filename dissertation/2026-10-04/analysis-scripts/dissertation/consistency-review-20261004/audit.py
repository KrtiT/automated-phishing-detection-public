"""Verify the bounded corrections against saved text and aggregate evidence."""

import csv
import json
import re
import subprocess
import sys
import unicodedata
from collections import Counter
from datetime import datetime, timezone
from zipfile import ZipFile

import fitz
from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation

import revise

sys.path.insert(0, str(revise.HERE.parent))
sys.path.insert(0, str(revise.CONTEXT / "gwu_advisor_update_source"))
import check_complete_package_20261002 as scientific
import check_followup_render_20261001 as render


def require(condition, message):
    if not condition:
        raise ValueError(message)


def rows(name):
    with (revise.BASE / "aggregate-data" / name).open() as stream:
        return list(csv.DictReader(stream))


def table_text(table):
    return [[cell.text for cell in row.cells] for row in table.rows]


def prose(document):
    return [paragraph.text for paragraph in document.paragraphs
            if not paragraph._p.xpath(".//w:instrText|.//w:hyperlink[@w:anchor]")]


def citations(markdown):
    body, tail = markdown.split("## References\n\n", 1)
    references = [text for text in tail.split("# Appendix A", 1)[0].strip().split("\n\n") if text]
    require(len(references) == 51, "Reference count")
    require(len(set(references)) == 51, "Duplicate reference")
    report = []
    used_spans = []
    for reference in references:
        authors, year = re.match(r"^(.*?) \(((?:19|20)\d{2}[a-z]?)\)", reference).groups()
        first = authors.split(",", 1)[0]
        names = re.findall(r"(?:^|, & |, )([A-ZŽ][^,]+), [A-Z]", authors)
        if len(names) > 2:
            stem = re.escape(first) + r" et al\."
        elif len(names) == 2:
            stem = re.escape(first) + r" (?:and|&) " + re.escape(names[1])
        else:
            stem = re.escape(first)
        pattern = re.compile(r"\b" + stem + r"(?:['’]s)?\s*(?:\(\s*|,\s*)(?P<years>(?:19|20)\d{2}[a-z]?"
                             r"(?:,\s*(?:19|20)\d{2}[a-z]?)*)(?=[;).,\s])")
        matches = [match for match in pattern.finditer(body) if year in re.findall(r"(?:19|20)\d{2}[a-z]?", match['years'])]
        require(bool(matches), "No exact author/year citation: " + reference)
        used_spans.extend((match.start(), match.end()) for match in matches)
        report.append({"reference": reference, "author_names": names, "year": year,
                       "citations": [{"line": body.count("\n", 0, match.start()) + 1,
                                      "text": match.group()} for match in matches]})
    year_tokens = list(re.finditer(r"(?<!\d)(?:19|20)\d{2}[a-z]?(?!\d)", body))
    residual = []
    for token in year_tokens:
        if not any(start <= token.start() and token.end() <= end for start, end in used_spans):
            line = body[body.rfind("\n", 0, token.start()) + 1:body.find("\n", token.end())]
            residual.append({"line": body.count("\n", 0, token.start()) + 1, "year": token.group(), "text": line})
    (revise.ROOT / "citation-audit.json").write_text(json.dumps({
        "references": report, "year_contexts_for_manual_review": residual,
        "scope": "Exact author/year identity and reverse year-context review, not full-source claim verification."
    }, indent=2, ensure_ascii=False) + "\n")
    return len(report), sum(len(item['citations']) for item in report), residual


def main():
    baseline = Document(revise.SOURCE)
    current = Document(revise.OUTPUT)
    markdown = revise.OUTPUT.with_suffix(".md").read_text()
    expected_md = revise.SOURCE.with_suffix(".md").read_text()
    for old, new, _ in revise.BODY_EDITS:
        expected_md = expected_md.replace(old, new, 1)
    expected_md = expected_md.replace("\n\nBreiman, L. (2001).", "\n\n" + revise.RFC + "\n\nBreiman, L. (2001).", 1)
    require(markdown == expected_md, "Unlogged Markdown change")
    expected = prose(baseline)
    for old, new, _ in revise.BODY_EDITS + revise.FRONT_EDITS:
        targets = [index for index, text in enumerate(expected) if old in text]
        require(len(targets) == 1, "Nonunique expected change")
        index = targets[0]
        expected[index] = expected[index].replace(old, new, 1)
    index = next(index for index, text in enumerate(expected) if text.startswith("Breiman, L. (2001)."))
    expected.insert(index, revise.RFC)
    require(prose(current) == expected, "Unlogged DOCX prose change or lost paragraph")
    require(current.paragraphs[2]._p.xml == baseline.paragraphs[2]._p.xml, "Credentials changed")
    require(len(current.tables) == 27, "Table inventory")
    require([table._tbl.xml for table in current.tables] == [table._tbl.xml for table in baseline.tables],
            "Table text, layout or styling changed")
    require(len(current.inline_shapes) == 4, "Figure inventory")
    with ZipFile(revise.SOURCE) as old, ZipFile(revise.OUTPUT) as new:
        for name in old.namelist():
            if name.startswith("word/media/"):
                require(old.read(name) == new.read(name), "Changed embedded figure")
    for section in current.sections:
        require((section.page_width.inches, section.page_height.inches) == (8.5, 11), "Page size")
        require((section.left_margin.inches, section.right_margin.inches) == (1.25, 1.25), "Side margins")
        require((section.top_margin.inches, section.bottom_margin.inches) == (1, 1), "Vertical margins")
    refs, citation_count, residual = citations(markdown)
    headings = set(re.findall(r"(?m)^#{2,3} (\d+\.\d+(?:\.\d+)?) ", markdown))
    targets = re.findall(r"\bSection (\d+\.\d+(?:\.\d+)?)", markdown)
    require(all(target in headings for target in targets), "Dangling section reference")
    placeholders = re.findall(r"(?i)\b(?:TODO|TBD|FIXME|lorem ipsum)\b|Error! Reference source not found|Error! Bookmark not defined", markdown)
    require(not placeholders, "Placeholder or broken reference")
    require(not current._element.xpath(".//w:ins|.//w:del"), "Unresolved tracked edits")
    for paragraph in current.paragraphs:
        if paragraph.text in markdown.split("## References", 1)[1].split("# Appendix A", 1)[0]:
            if re.match(r"^[A-Z].*\((?:19|20)\d{2}", paragraph.text):
                require(paragraph.paragraph_format.left_indent.inches == .5 and
                        paragraph.paragraph_format.first_line_indent.inches == -.5, "Reference hanging indent")
                require(any(run.italic for run in paragraph.runs), "Missing reference title/container emphasis")
    gates = rows("primary-gates.csv")
    primary_rows = table_text(current.tables[16])[1:]
    require(len(gates) == len(primary_rows) == 22, "Original hypothesis checks")
    for saved, displayed in zip(gates, primary_rows):
        unit = " ms" if saved['name'] == "http_pooled_p95_ms" else "%"
        if '_minus_' in saved['name']:
            unit = " pp"
        scale = 1 if unit == " ms" else 100
        operand = f"{float(saved['estimate']) * scale:.4f}{unit}"
        threshold = f"{saved['operator']} {float(saved['threshold']) * scale:.4f}{unit}"
        if unit == " ms":
            threshold = f"{saved['operator']} {float(saved['threshold']):g}{unit}"
        require(displayed == [saved['hypothesis'], saved['name'].replace('_', ' '), operand, threshold, saved['status']],
                "Original gate differs from saved aggregate: " + saved['name'])
    groups = rows("operational-groups.csv")
    require(len(groups) == 25 and len(rows("operational-runs.csv")) == 125, "Original operational inventory")
    for saved, displayed in zip(groups, table_text(current.tables[8])[1:]):
        expected_values = [f"{float(saved[field]):.3f}" for field in ['p50_ms', 'p95_ms', 'p99_ms']]
        expected_values += [f"{saved['request_errors']}/{saved['request_count']}",
                            f"{100 * float(saved['physical_invocation_fraction']):.4f}%"]
        require(displayed[1:] == expected_values, "Original operational values differ")
    numeric = scientific.check_numeric_tables(current)
    deck = Presentation(next((revise.ROOT / "advisor").glob("*.pptx")))
    deck_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    questions = set(re.findall(r"(?m)^RQ[123]: .+$", markdown))
    require(len(questions) == 3 and all(question in deck_text for question in questions), "RQ deck mismatch")
    require(len(deck.slides) == 25, "Deck inventory")
    for value in ['0.20325', '79.68%', '99,999', '799 / 100,000', '8,622', '251', '184']:
        require(value in deck_text, "Missing deck result: " + value)
    for path in (revise.ROOT / 'advisor').iterdir():
        require(revise.digest(path) == revise.digest(revise.BASE / 'advisor' / path.name), "Deck changed")
    preserved = 0
    for line in (revise.BASE / 'SHA256SUMS.txt').read_text().splitlines():
        expected_hash, name = line.split('  ', 1)
        require(revise.digest(revise.BASE / name) == expected_hash, "Changed sealed member: " + name)
        preserved += 1
    before = json.loads((revise.HERE / 'preservation-before.json').read_text())
    require(revise.digest(revise.OPEN_COPY) == before['sha256'], "Open author copy changed on disk")
    frozen = revise.CONTEXT / 'gwu_working/study-followup-development-20261001'
    revision = subprocess.check_output(['git', '-C', str(frozen), 'rev-parse', 'HEAD'], text=True).strip()
    dirty = subprocess.check_output(['git', '-C', str(frozen), 'status', '--porcelain'], text=True)
    require(revision == 'ef8ba5f0b357cf3dd60c4d663e6297d13334460c' and not dirty, "Frozen checkout changed")
    render.HERE = render.REVIEW = revise.MANUSCRIPT
    render.PDF = revise.OUTPUT.with_suffix('.pdf')
    render.main()
    rendered = fitz.open(render.PDF)
    full_text = scientific.normalize('\n'.join(page.get_text(clip=fitz.Rect(85, 65, 528, 725)) for page in rendered))
    fragments = 0
    for paragraph in current._element.xpath('.//w:p'):
        if paragraph.xpath('.//w:instrText|.//w:hyperlink[@w:anchor]'):
            continue
        text = ''.join(paragraph.xpath('.//w:t/text()'))
        if text.strip():
            require(scientific.normalize(text) in full_text, 'Missing PDF text: ' + text[:100])
            fragments += 1
    report = {
        'status': 'bounded_consistency_checks_passed', 'checked_at': datetime.now(timezone.utc).isoformat(),
        'logged_content_changes': 9, 'logged_formatting_corrections': 1,
        'references': refs, 'matched_author_year_citation_links': citation_count,
        'reverse_year_contexts_requiring_manual_review': len(residual), 'section_cross_references': len(targets),
        'unchanged_tables': 27, 'unchanged_figures': 4, 'primary_checks': 22,
        'original_check_decisions': dict(Counter(row['status'] for row in gates)),
        'operational_groups': 25, 'operational_cells': 125, 'saved_D_S_checks': numeric,
        'advisor_slides': len(deck.slides), 'exact_RQ_deck_matches': 3,
        'sealed_members_preserved': preserved, 'open_author_copy_sha256': before['sha256'],
        'frozen_revision': revision, 'frozen_checkout_clean': True,
        'pdf_pages': len(rendered), 'rendered_paragraphs_and_cells': fragments,
        'files': {path.name: revise.digest(path) for path in [revise.OUTPUT, revise.OUTPUT.with_suffix('.md'), render.PDF]},
        'limits': ['Not institutional committee certification or AIR/Turnitin clearance.',
                   'Exact citation identity is not full-source verification of every scholarly claim.',
                   'Manual visual and reverse-year reviews are recorded separately.'],
    }
    (revise.ROOT / 'verification.json').write_text(json.dumps(report, indent=2) + '\n')
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    main()
