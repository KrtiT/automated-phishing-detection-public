"""Add verified source support to a separate manuscript review copy."""

import hashlib
import importlib.util
import json
import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


HERE = Path(__file__).resolve().parent
CONTEXT = HERE.parent.parent
AUTHOR = CONTEXT / "deliverables/Tallam_Dissertation_Author_Review_2026-10-04"
BASE = AUTHOR / "Consistency_Reviewed_2026-10-04"
ROOT = AUTHOR / "Citations_Expanded_2026-10-04"
MANUSCRIPT = ROOT / "manuscript"
SOURCE = BASE / "manuscript/Tallam_Krti_Praxis_Consistency_Reviewed_2026-10-04.docx"
OUTPUT = MANUSCRIPT / "Tallam_Krti_Praxis_Citations_Expanded_2026-10-04.docx"
OPEN_COPY = AUTHOR / "Tallam_Krti_Praxis_Author_Review_2026-10-04.docx"

spec = importlib.util.spec_from_file_location("prior_revision", HERE.parent / "consistency-review-20261004/revise.py")
prior = importlib.util.module_from_spec(spec)
spec.loader.exec_module(prior)

EDITS = [
    {
        "section": "2.3", "source": "tranco",
        "old": "Tranco records are a popularity-based control rather than verified benign evidence.",
        "new": "Tranco records are a popularity-based control rather than verified benign evidence. Le Pochat et al. (2019) show how the composition, stability and susceptibility to manipulation of popularity rankings can affect samples used in security research. Their findings support keeping popularity-based controls separate from independently verified negative labels.",
    },
    {
        "section": "2.4", "source": "axelsson",
        "old": "An inline alerting system can be unusable even with high aggregate accuracy if its false-positive rate is too large.",
        "new": "An inline alerting system can be unusable even with high aggregate accuracy if its false-positive rate is too large. Axelsson's (2000) base-rate analysis of intrusion detection explains why rare events can make false alarms dominate the alert population despite high detection sensitivity. That analysis motivates examining low-FPR behavior, but does not prescribe the particular 1% ceiling adopted here.",
    },
    {
        "section": "2.6", "source": "dempster",
        "old": "A Gaussian mixture model (GMM) represents the development feature distribution as a weighted combination of Gaussian components.",
        "new": "A Gaussian mixture model (GMM) represents the development feature distribution as a weighted combination of Gaussian components. In the expectation–maximization formulation, component membership is unobserved; estimation alternates between conditional membership weights and parameter updates (Dempster et al., 1977).",
    },
    {
        "section": "2.9", "source": "tail",
        "old": "Connection management, request scheduling, validation and response processing can affect observed latency even when model parameters are unchanged.",
        "new": "Connection management, request scheduling, validation and response processing can affect observed latency even when model parameters are unchanged. Dean and Barroso (2013) explain why latency variability matters for the responsiveness of large interactive services. Their scale differs from this local gateway, but the distinction supports examining the slow tail rather than average service time alone.",
    },
    {
        "section": "3.6", "source": "adamw",
        "old": "One AdamW parameter group contains every trainable parameter",
        "new": "Optimization uses AdamW's decoupled weight-decay formulation (Loshchilov & Hutter, 2019). One AdamW parameter group contains every trainable parameter",
    },
    {
        "section": "3.9", "source": "existing_schwarz",
        "old": "The model with minimum training Bayesian information criterion (BIC) is selected;",
        "new": "The model with minimum training Bayesian information criterion (BIC; Schwarz, 1978) is selected;",
    },
    {
        "section": "3.12", "source": "schroeder",
        "old": "Closed-loop workers time requests from submission through body reading and response validation, excluding unsent local backlog; there are no retries.",
        "new": "Closed-loop workers time requests from submission through body reading and response validation, excluding unsent local backlog; there are no retries. This is a closed workload in the sense of Schroeder et al. (2006): each worker submits another request only after its preceding request terminates.",
    },
    {
        "section": "5.3", "source": "tranco",
        "old": "Tranco controls provide an additional alarm about generality, but popularity must not be relabeled as verified benign truth.",
        "new": "Tranco controls provide an additional alarm about generality, but popularity must not be relabeled as verified benign truth (Le Pochat et al., 2019).",
    },
]

ADDITIONS = [
    {
        "section": "2.4", "source": "lipton", "before": "## 2.5 Selective Cascades",
        "text": "Prevalence projections require a further distinction between changing class proportions and changing class-conditional behavior. Lipton et al. (2018) study label shift, in which P(Y) changes while P(X|Y) remains fixed, and develop a method for estimating the changed class proportions. The prevalence projections reported here hold the measured class-conditional rates fixed and vary an assumed phishing proportion; they neither estimate a deployment prevalence nor implement that paper's shift-correction method. Their value is to show how the alert burden depends on the assumed class mixture, separately from the measured source-transfer problem.",
    },
    {
        "section": "2.9", "source": "schroeder", "before": "Service correctness also has multiple layers.",
        "text": "The workload generator also determines what a throughput or latency result means. Schroeder et al. (2006) distinguish closed workloads, in which completions govern subsequent submissions, from open workloads with arrivals independent of completion. Their experiments show that these choices can produce materially different response times and scheduling behavior. The present fixed-concurrency replays therefore characterize the tested closed-loop service path. They support a controlled comparison of client designs, rather than a capacity claim for an externally imposed production arrival rate.",
    },
]

REFERENCES = [
    {
        "source": "axelsson", "before": "Basit, A.",
        "prefix": "Axelsson, S. (2000). The base-rate fallacy and the difficulty of intrusion detection. ",
        "italic": "ACM Transactions on Information and System Security, 3",
        "suffix": "(3), 186–205. https://doi.org/10.1145/357830.357849",
    },
    {
        "source": "tail", "before": "Efron, B.",
        "prefix": "Dean, J., & Barroso, L. A. (2013). The tail at scale. ",
        "italic": "Communications of the ACM, 56",
        "suffix": "(2), 74–80. https://doi.org/10.1145/2408776.2408794",
    },
    {
        "source": "dempster", "before": "Efron, B.",
        "prefix": "Dempster, A. P., Laird, N. M., & Rubin, D. B. (1977). Maximum likelihood from incomplete data via the EM algorithm. ",
        "italic": "Journal of the Royal Statistical Society: Series B (Methodological), 39",
        "suffix": "(1), 1–22. https://doi.org/10.1111/j.2517-6161.1977.tb01600.x",
    },
    {
        "source": "tranco", "before": "Li, L.",
        "prefix": "Le Pochat, V., Van Goethem, T., Tajalizadehkhoob, S., Korczyński, M., & Joosen, W. (2019). Tranco: A research-oriented top sites ranking hardened against manipulation. ",
        "italic": "Network and Distributed System Security Symposium",
        "suffix": ". https://doi.org/10.14722/ndss.2019.23386",
    },
    {
        "source": "lipton", "before": "Lu, J.",
        "prefix": "Lipton, Z. C., Wang, Y.-X., & Smola, A. J. (2018). Detecting and correcting for label shift with black box predictors. ",
        "italic": "Proceedings of Machine Learning Research, 80",
        "suffix": ", 3122–3130. https://proceedings.mlr.press/v80/lipton18a.html",
    },
    {
        "source": "adamw", "before": "Lu, J.",
        "prefix": "Loshchilov, I., & Hutter, F. (2019). Decoupled weight decay regularization. ",
        "italic": "International Conference on Learning Representations",
        "suffix": ". https://arxiv.org/abs/1711.05101",
    },
    {
        "source": "schroeder", "before": "Schwarz, G.",
        "prefix": "Schroeder, B., Wierman, A., & Harchol-Balter, M. (2006). Open versus closed: A cautionary tale. ",
        "italic": "3rd USENIX Symposium on Networked Systems Design & Implementation",
        "suffix": ", 239–252. https://www.usenix.org/legacy/events/nsdi06/tech/schroeder.html",
    },
]


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def target(document, prefix):
    matches = [paragraph for paragraph in document.paragraphs
               if paragraph.text.startswith(prefix)
               and not paragraph._p.xpath(".//w:instrText|.//w:hyperlink[@w:anchor]")]
    if len(matches) != 1:
        raise ValueError("Ambiguous paragraph target: " + prefix)
    return matches[0]


def build():
    if ROOT.exists():
        raise FileExistsError("Refusing to overwrite an existing review package")
    if digest(SOURCE) != "0606ca3370e903c66ce39a503663f373c8cc2bb76c3a7fd0acecdd4072620a66":
        raise ValueError("The reviewed manuscript changed; reconcile before editing")
    preserved = {str(path): digest(path) for path in [OPEN_COPY, SOURCE, SOURCE.with_suffix(".pdf"), SOURCE.with_suffix(".md")]}
    MANUSCRIPT.mkdir(parents=True)
    document = Document(SOURCE)
    markdown = SOURCE.with_suffix(".md").read_text()
    ledger = []
    for edit in EDITS:
        matches = [paragraph for paragraph in document.paragraphs if edit["old"] in paragraph.text]
        if len(matches) != 1 or markdown.count(edit["old"]) != 1:
            raise ValueError("Ambiguous span: " + edit["old"])
        prior.replace_span(matches[0], edit["old"], edit["new"])
        markdown = markdown.replace(edit["old"], edit["new"], 1)
        ledger.append({"kind": "replace", **edit})
    for addition in ADDITIONS:
        prefix = addition["before"].removeprefix("## ")
        anchor = target(document, prefix)
        element = deepcopy(target(document, "An inference endpoint includes more than a classifier.")._p)
        anchor._p.addprevious(element)
        paragraph = Paragraph(element, anchor._parent)
        paragraph.clear()
        paragraph.add_run(addition["text"])
        if markdown.count(addition["before"]) != 1:
            raise ValueError("Ambiguous insertion anchor")
        markdown = markdown.replace(addition["before"], addition["text"] + "\n\n" + addition["before"], 1)
        ledger.append({"kind": "insert", **addition})
    for reference in REFERENCES:
        anchor = target(document, reference["before"])
        element = deepcopy(anchor._p)
        anchor._p.addprevious(element)
        paragraph = Paragraph(element, anchor._parent)
        paragraph.clear()
        paragraph.add_run(reference["prefix"])
        paragraph.add_run(reference["italic"]).italic = True
        paragraph.add_run(reference["suffix"])
        markdown = markdown.replace("\n\n" + reference["before"], "\n\n" + paragraph.text + "\n\n" + reference["before"], 1)
        ledger.append({"kind": "reference", "source": reference["source"], "text": paragraph.text})
    document.save(OUTPUT)
    OUTPUT.with_suffix(".md").write_text(markdown)
    for path in (BASE / "manuscript").iterdir():
        if path.suffix in {".png", ".pdf"} and not path.name.startswith(("Tallam_", "contact-")):
            shutil.copy2(path, MANUSCRIPT / path.name)
    shutil.copy2(BASE / "manuscript/gwu-navigation-20261001.json", MANUSCRIPT)
    shutil.copytree(BASE / "advisor", ROOT / "advisor")
    (ROOT / "changes.json").write_text(json.dumps(ledger, indent=2, ensure_ascii=False) + "\n")
    (HERE / "preservation-before.json").write_text(json.dumps(preserved, indent=2) + "\n")
    print(json.dumps({"output": str(OUTPUT), "new_references": len(REFERENCES), "new_paragraphs": len(ADDITIONS), "targeted_edits": len(EDITS)}))


def refresh():
    prior.MANUSCRIPT = MANUSCRIPT
    prior.OUTPUT = OUTPUT
    prior.refresh()


if __name__ == "__main__":
    {"build": build, "refresh": refresh}[sys.argv[1]]()
