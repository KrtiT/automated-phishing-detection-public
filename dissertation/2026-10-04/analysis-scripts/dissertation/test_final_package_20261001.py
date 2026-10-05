import importlib.util
import json
import unittest
from tempfile import TemporaryDirectory
from pathlib import Path

HERE = Path(__file__).resolve().parent
SPEC = importlib.util.spec_from_file_location("manuscript_builder", HERE / "build_v3_working_manuscript.py")
BUILDER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILDER)


class FinalPackageTests(unittest.TestCase):
    def test_advisor_cover_styles_every_line_below_the_logo(self):
        from pptx import Presentation
        from pptx.util import Inches

        spec = importlib.util.spec_from_file_location("final_deck", HERE.parent / "gwu_advisor_update_source/build_advisor_results_20261001.py")
        deck = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(deck)
        with TemporaryDirectory() as temporary:
            deck.OUTPUT = Path(temporary) / "results.pptx"
            deck.main()
            presentation = Presentation(deck.OUTPUT)
        cover = presentation.slides[0]
        title = next(shape for shape in cover.shapes if shape.has_text_frame and "Automated Phishing Detection" in shape.text)
        self.assertGreaterEqual(title.top, Inches(1.6))
        self.assertLessEqual(title.width, Inches(5.2))
        for shape in cover.shapes:
            if not shape.has_text_frame or not shape.text.strip():
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    self.assertEqual(run.font.name, "Arial")
                    self.assertEqual(run.font.color.rgb, deck.LAYOUT.WHITE)

    def test_large_tables_can_paginate_without_chaining_every_row(self):
        table = BUILDER._table([["Measure", "Value"]] + [[f"Row {index}", str(index)] for index in range(30)])
        rows = table.findall(f"./{BUILDER.W}tr")
        self.assertIsNotNone(rows[0].find(f".//{BUILDER.W}keepNext"))
        for row in rows[1:]:
            self.assertIsNone(row.find(f".//{BUILDER.W}keepNext"))
            self.assertIsNotNone(row.find(f"./{BUILDER.W}trPr/{BUILDER.W}cantSplit"))

    def test_all_primary_checks_are_complete_and_preserved_as_adverse(self):
        result = json.loads((HERE / "final-evidence-20261001/primary-results.json").read_text())
        for hypothesis in result["primary"]["hypotheses"].values():
            self.assertTrue(hypothesis["complete"])
            self.assertEqual(hypothesis["decision"], "not_supported")
        body = (HERE / "Tallam_Krti_Praxis_v3_Body_Results_2026-10-01.md").read_text()
        self.assertEqual(body.count("# Chapter "), 5)
        for text in ["all 22 primary checks are complete", "72 eligible complete cells", "remaining 53", "364.3004101 ms", "90.5487%", "H1 is not supported", "H2 is not supported", "H3 is not supported", "digests of the label vectors actually consumed", "not a single uninterrupted experiment"]:
            self.assertIn(text, body)
        for text in ["H1 and H3 await", "No new research measurements have restarted", "No continuation measurements have restarted", "Historical eligibility for checkpointed recovery is not yet"]:
            self.assertNotIn(text, body)

    def test_scope_and_secondary_coverage_match_the_frozen_schema(self):
        result = json.loads((HERE / "final-evidence-20261001/secondary-verification.json").read_text())
        self.assertEqual(result["population_detector_rows"], 153)
        self.assertEqual(result["labeled_metric_rows"], 43)
        self.assertEqual(result["calibration_bins"], 430)
        self.assertEqual(result["projection_rows"], 129)
        self.assertEqual(result["monitor_windows"], 396)
        self.assertEqual(result["psi_feature_scores"], 3432)


if __name__ == "__main__":
    unittest.main()
