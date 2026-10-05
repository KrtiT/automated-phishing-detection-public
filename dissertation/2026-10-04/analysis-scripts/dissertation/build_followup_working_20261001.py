"""Build the current manuscript using the supplied GWU template; preserve originals."""

import hashlib
import json
import re
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

import build_v3_working_manuscript as builder
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement, parse_xml
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from lxml import etree

SOURCE = Path(__file__).resolve().parent
HERE = SOURCE / "followup-20261001/manuscript-work"
BODY = HERE / "Tallam_Krti_Praxis_Engineering_Working_2026-10-01.md"
OUTPUT = HERE / "Tallam_Krti_Praxis_Engineering_Working_2026-10-01.docx"
TEMPLATE = SOURCE.parent / "attachments/pduS8C/DEng Praxis Template - Online Programs Rev 2026.docx"
ORIGINAL = SOURCE.parent / "attachments/foUeB5/Tallam_Krti_Praxis_2026-06-15.docx"
TITLE = "Automated Phishing Detection for Frontier AI Inference"
FIGURES = {
    "!Frozen inference and evaluation dataflow": SOURCE / "gwu-system-dataflow-20261001.png",
    "!Follow-up source partitions": SOURCE / "followup-20261001/figures/followup-source-partitions.png",
    "!Follow-up paired recall": SOURCE / "followup-20261001/figures/followup-paired-recall.png",
    "!Follow-up calibration": SOURCE / "followup-20261001/figures/followup-calibration.png",
}
REFERENCE_ITALICS = [
    "Frontiers in Computer Science, 8", "Electronics, 15", "Biometrika, 26",
    "The Annals of Statistics, 7", "Scandinavian Journal of Statistics, 6",
    "Expert Systems with Applications, 333",
    "Findings of the Association for Computational Linguistics: EMNLP 2021",
    "Psychometrika, 12", "Computers & Security, 136", "Computer Networks, 245",
    "The Annals of Statistics, 6",
    "Journal of the Royal Statistical Society: Series B (Methodological), 58",
    "Proceedings of the AAAI Conference on Artificial Intelligence, 38",
    "Data in Brief, 68", "30th USENIX Security Symposium",
    "31st USENIX Security Symposium", "28th USENIX Security Symposium",
    "Proceedings of Machine Learning Research, 70", "Proceedings of Machine Learning Research, 97",
    "Engineering Applications of Artificial Intelligence, 104", "Web page phishing detection (Version 3)",
    "PhiUSIIL Phishing URL (Website)",
    "PhishVN: A time-stamped Vietnamese URL phishing dataset with impersonation-scenario labels and confidence tiers (Version 4)",
]
ABSTRACT = """This study develops and evaluates a raw-URL inference gateway through three questions: representation value, Gaussian-mixture-guided escalation under source/domain shift, and inline detection–service tradeoffs. The initial investigation includes 34,593 internal records, 8,701 external records, 125 real-HTTP cells and 22 primary checks. Nine component checks pass. Structural features improve internal recall by 64.07 percentage points over length alone, the monitor detects 87.12% of external windows, and primary requests complete without errors. External specificity, useful escalation and joint service requirements remain unmet; H1–H3 are not supported under unchanged rules. Mechanism-directed diagnosis leads to two bounded comparisons. A transport-neutral structural model is fitted and calibrated on the original development partitions and compared with the unchanged model on 8,622 eligible URLs from an additional 2020 benchmark. Exact feature and score invariance holds under every prescribed HTTP/HTTPS companion transformation. At frozen thresholds, the candidate produces 251 fewer false positives and 184 fewer detected positives. Recall difference is −3.96 percentage points (97.5% domain-cluster interval [−5.05, −3.00]); false-positive rate remains 92.90%, so detection requirement D is not met. The contribution connects a working artifact, diagnosed mechanisms and measured modification tradeoffs, separating representation stability from external risk and HTTP cost. Publisher labels and retrospective comparisons bound interpretation. Controlled client-topology measurements remain pending; this working manuscript is not the completed extension."""


def field(paragraph, instruction, cached="1"):
    run = paragraph.add_run()._r
    start = OxmlElement("w:fldChar")
    start.set(builder.W + "fldCharType", "begin")
    code = OxmlElement("w:instrText")
    code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(builder.W + "fldCharType", "separate")
    value = OxmlElement("w:t")
    value.text = cached
    end = OxmlElement("w:fldChar")
    end.set(builder.W + "fldCharType", "end")
    for node in (start, code, separate, value, end):
        run.append(node)


def bookmark(paragraph, name, identifier):
    start = OxmlElement("w:bookmarkStart")
    start.set(builder.W + "id", str(identifier))
    start.set(builder.W + "name", name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(builder.W + "id", str(identifier))
    paragraph._p.append(start)
    paragraph._p.append(end)


def front_heading(document, text, name, entries, *, include=True):
    paragraph = document.add_paragraph(text)
    paragraph.alignment = 1
    paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(18)
    paragraph.runs[0].bold = True
    bookmark(paragraph, name, 1000 + len(document.paragraphs))
    if include:
        entries.append({"text": text, "bookmark": name, "level": 1, "kind": "front"})
    return paragraph


def index_entry(document, entry, pages):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.left_indent = Inches(.18 * (entry["level"] - 1))
    paragraph.paragraph_format.right_indent = Inches(.3)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run(entry["text"] + "\t")
    if entry["kind"] == "front":
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(builder.W + "anchor", entry["bookmark"])
        hyperlink.append(builder._run(pages.get(entry["bookmark"], "—")))
        paragraph._p.append(hyperlink)
    else:
        field(paragraph, f' PAGEREF {entry["bookmark"]} \\h ', pages.get(entry["bookmark"], "—"))


def current_front(document, body_entries, table_entries, figure_entries, pages):
    with ZipFile(ORIGINAL) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    _, original_elements = builder._validate_source(ORIGINAL.read_bytes(), root)
    def original_text(index):
        return "".join(original_elements[index].itertext())

    entries = []
    title = document.add_paragraph(TITLE)
    title.alignment = 1
    title.runs[0].bold = True
    title.paragraph_format.space_after = Pt(30)
    for text, gap in [("by Krti Tallam", 24), ("\n".join(original_text(index) for index in range(6, 10)), 24), ("A Praxis submitted to\n\nThe Faculty of\nThe School of Engineering and Applied Science\nof The George Washington University\nin partial fulfillment of the requirements\nfor the degree of Doctor of Engineering", 24), ("October 1, 2026", 24), ("Praxis directed by\n\nMazen Mheish\nProfessorial Lecturer in Engineering and Applied Science", 0)]:
        paragraph = document.add_paragraph(text)
        paragraph.alignment = 1
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.space_after = Pt(gap)
    front_heading(document, "Praxis Research Committee", "committee", entries, include=False)
    document.add_paragraph(TITLE).alignment = 1
    document.add_paragraph("Krti Tallam").alignment = 1
    for index in (54, 55, 56):
        paragraph = document.add_paragraph(original_text(index))
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.space_after = Pt(18)
    document.add_paragraph("Committee information is carried forward from the author-supplied manuscript. University certification of the final examination and institutional approval is not asserted in this copy.")
    spacer = document.add_paragraph()
    spacer.paragraph_format.page_break_before = True
    spacer.paragraph_format.keep_with_next = True
    copyright = front_heading(document, original_text(67), "copyright", entries, include=False)
    copyright.paragraph_format.page_break_before = False
    copyright.paragraph_format.space_before = Pt(280)
    document.add_paragraph(original_text(68)).alignment = 1
    front_heading(document, "Dedication", "dedication", entries)
    document.add_paragraph(original_text(71))
    front_heading(document, "Acknowledgements", "acknowledgements", entries)
    for index in (74, 75):
        document.add_paragraph(original_text(index))
    front_heading(document, "Abstract of Praxis", "abstract", entries)
    document.add_paragraph(TITLE).alignment = 1
    document.add_paragraph(ABSTRACT)
    front_heading(document, "Table of Contents", "contents", entries, include=False)
    entries.extend([
        {"text": "List of Figures", "bookmark": "list_figures", "level": 1, "kind": "front"},
        {"text": "List of Tables", "bookmark": "list_tables", "level": 1, "kind": "front"},
        {"text": "List of Symbols", "bookmark": "symbols", "level": 1, "kind": "front"},
        {"text": "List of Acronyms", "bookmark": "acronyms", "level": 1, "kind": "front"},
    ])
    for entry in entries + body_entries:
        index_entry(document, entry, pages)
    front_heading(document, "List of Figures", "list_figures", entries, include=False)
    for entry in figure_entries:
        index_entry(document, entry, pages)
    front_heading(document, "List of Tables", "list_tables", entries, include=False)
    for entry in table_entries:
        index_entry(document, entry, pages)
    front_heading(document, "List of Symbols", "symbols", entries, include=False)
    definitions = [
        ("TP, FP, TN, FN", "True-positive, false-positive, true-negative and false-negative counts against source reference labels."),
        ("P, N", "Positive and negative stratum sizes in confusion-count tables; a total stream size is identified explicitly where used."),
        ("X; P(X)", "Observed feature vector; its marginal distribution."),
        ("τ", "Validation-selected decision threshold, unchanged during evaluation."),
        ("p50, p95, p99", "50th, 95th and 99th percentiles of individual measured terminal request latencies."),
        ("pp", "Percentage points; the unit of differences between proportions expressed as percentages."),
        ("c", "HTTP client concurrency, such as c64 for 64 concurrent clients."),
    ]
    for term, meaning in definitions:
        paragraph = document.add_paragraph(f"{term} — {meaning}")
        paragraph.paragraph_format.space_after = Pt(6)
    front_heading(document, "List of Acronyms", "acronyms", entries, include=False)
    acronyms = {
        "AI": "Artificial intelligence", "AP": "Average precision", "API": "Application programming interface", "ASGI": "Asynchronous Server Gateway Interface", "AUC": "Area under the curve", "BIC": "Bayesian information criterion", "CP": "Clopper–Pearson", "CPU": "Central processing unit", "CSV": "Comma-separated values", "DNS": "Domain Name System", "ECE": "Expected calibration error", "FNR": "False-negative rate", "FPR": "False-positive rate", "GMM": "Gaussian mixture model", "HTTP": "Hypertext Transfer Protocol", "HTTPS": "Hypertext Transfer Protocol Secure", "JSON": "JavaScript Object Notation", "MCC": "Matthews correlation coefficient", "MMD": "Maximum mean discrepancy", "NCSC": "National Cyber Security Center (source designation in PhishVN)", "PSI": "Population stability index", "RF": "Random Forest", "ROC": "Receiver operating characteristic", "TPR": "True-positive rate", "URL": "Uniform resource locator"
    }
    for term, meaning in acronyms.items():
        paragraph = document.add_paragraph(f"{term} — {meaning}")
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.space_after = Pt(9)
    return entries


def main():
    body = BODY.read_text()
    document = Document(TEMPLATE)
    if "Caption" not in document.styles:
        document.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH).base_style = document.styles["Normal"]
    for element in list(document._element.body)[:-1]:
        document._element.body.remove(element)
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1.25)
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.header_distance = section.footer_distance = Inches(.5)
    section.different_first_page_header_footer = True
    numbering = section._sectPr.find(builder.W + "pgNumType")
    if numbering is None:
        numbering = OxmlElement("w:pgNumType")
        section._sectPr.append(numbering)
    numbering.set(builder.W + "fmt", "lowerRoman")
    numbering.set(builder.W + "start", "1")
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 2
        style.paragraph_format.space_after = Pt(0)
    document.core_properties.author = "Krti Tallam"
    document.core_properties.title = TITLE
    document.core_properties.subject = "GWU D.Eng. Praxis — engineering working copy; service comparison pending"
    document.core_properties.comments = "Current scholarly text; not institutional approval or certification."
    document.core_properties.last_modified_by = "Krti Tallam"
    elements = [parse_xml(etree.tostring(element)) for element in builder._body_elements(body)]
    body_entries, table_entries, figure_entries = [], [], []
    in_references = False
    for index, element in enumerate(elements):
        if element.tag != builder.W + "p":
            continue
        paragraph = Paragraph(element, document._body)
        if paragraph.style.name.startswith("Heading"):
            in_references = paragraph.text == "References"
            level = 1 if paragraph.text == "References" else int(paragraph.style.name[-1])
            if paragraph.text == "References":
                paragraph.style = "Heading 1"
                paragraph.paragraph_format.page_break_before = True
                paragraph.alignment = 1
            name = f"heading_{len(body_entries) + 1}"
            bookmark(paragraph, name, 2000 + len(body_entries))
            body_entries.append({"text": paragraph.text, "bookmark": name, "level": level, "kind": "body"})
        elif in_references:
            reference = paragraph.text
            for italic_text in REFERENCE_ITALICS:
                if italic_text in reference:
                    before, after = reference.split(italic_text, 1)
                    paragraph.clear()
                    paragraph.add_run(before)
                    paragraph.add_run(italic_text).italic = True
                    paragraph.add_run(after)
                    break
        if paragraph.text in FIGURES:
            figure_path = FIGURES[paragraph.text]
            paragraph.clear()
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(figure_path), width=Inches(6))
        if re.match(r"Figure [\dA-Z]+\.\d+\. ", paragraph.text):
            paragraph.style = "Caption"
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.space_after = Pt(12)
            name = f"figure_{len(figure_entries) + 1}"
            bookmark(paragraph, name, 4000 + len(figure_entries))
            figure_entries.append({"text": paragraph.text, "bookmark": name, "level": 1, "kind": "figure"})
        if re.match(r"Table [\dA-Z]+\.\d+[a-z]?\. ", paragraph.text) and index + 1 < len(elements) and elements[index + 1].tag == builder.W + "tbl":
            paragraph.style = "Caption"
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.font.size = Pt(12)
                run.italic = False
            name = f"table_{len(table_entries) + 1}"
            bookmark(paragraph, name, 3000 + len(table_entries))
            table_entries.append({"text": paragraph.text, "bookmark": name, "level": 1, "kind": "table"})
        if paragraph.text.startswith("CyberSentinel will test low-FPR") or paragraph.text.startswith("Published phishing-URL studies often optimize"):
            for run in paragraph.runs:
                run.italic = True
    page_map = HERE / "gwu-pagination-map-20261001.json"
    pages = json.loads(page_map.read_text()) if page_map.exists() else {}
    front_entries = current_front(document, body_entries, table_entries, figure_entries, pages)
    body_section = document.add_section(WD_SECTION_START.NEW_PAGE)
    body_section.different_first_page_header_footer = False
    body_section.footer.is_linked_to_previous = False
    body_numbering = body_section._sectPr.find(builder.W + "pgNumType")
    body_numbering.set(builder.W + "fmt", "decimal")
    body_numbering.set(builder.W + "start", "1")
    body_element = document._element.body
    for index, element in enumerate(elements):
        if index == 0:
            page_break = element.find(".//" + builder.W + "pageBreakBefore")
            if page_break is not None:
                page_break.getparent().remove(page_break)
        if element.tag == builder.W + "tbl":
            properties = element.find(builder.W + "tblPr")
            layout = OxmlElement("w:tblLayout")
            layout.set(builder.W + "type", "autofit")
            properties.append(layout)
            width = properties.find(builder.W + "tblW")
            width.set(builder.W + "w", "8640")
            width.set(builder.W + "type", "dxa")
        body_element.insert(len(body_element) - 1, deepcopy(element))
    for current_section in document.sections:
        footer = current_section.footer.paragraphs[0]
        footer.clear()
        footer.alignment = 1
        field(footer, " PAGE ")
    document.sections[0].first_page_footer.paragraphs[0].clear()
    document.save(OUTPUT)
    (HERE / "gwu-navigation-20261001.json").write_text(json.dumps(front_entries + body_entries + table_entries + figure_entries, indent=2) + "\n")
    (HERE / "Current_Abstract_2026-10-01.txt").write_text(ABSTRACT + "\n")
    metadata = {"title": document.core_properties.title, "author": "Krti Tallam", "date": "October 1, 2026", "abstract": ABSTRACT, "lang": "en-US"}
    (HERE / "results-pdf-metadata-20261001.json").write_text(json.dumps(metadata, indent=2) + "\n")
    print(OUTPUT)
    print("sha256=" + hashlib.sha256(OUTPUT.read_bytes()).hexdigest())


if __name__ == "__main__":
    main()
