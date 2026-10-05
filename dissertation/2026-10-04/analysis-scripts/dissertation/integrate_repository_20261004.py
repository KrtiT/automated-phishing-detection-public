"""Build the bounded repository-linked edition from the sealed public commit.

Usage: python integrate.py REPOSITORY NEW_OUTPUT_DIRECTORY
Requires python-docx and lxml. PDF rendering and verification are separate steps.
No research observations, model execution or network requests are involved.
"""

import hashlib
import json
import posixpath
import subprocess
import sys
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.text.paragraph import Paragraph
from lxml import etree

BASE_COMMIT = "91995dd3fa6f0661d185999bab90a6fabb25d962"
PACKAGE_PATH = "dissertation/2026-10-04"
RELEASE = "https://github.com/KrtiT/automated-phishing-detection-public/releases/tag/research-record-2026-10-04"
TAG_BASE = "https://github.com/KrtiT/automated-phishing-detection-public/tree/research-record-2026-10-04/"
TITLE = (
    "Automated phishing detection for frontier AI inference: Retained research record"
)
REFERENCE = f"Tallam, K. (2026). {TITLE} (research-record-2026-10-04) [Computer software and data set]. GitHub. {RELEASE}"
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

EDITS = [
    (
        "The linked public evidence records collectively identify",
        "The versioned public research record is cited as Tallam (2026); Appendix A connects it to the manuscript and the retained observations. The linked public evidence records collectively identify",
    ),
    ("private model-artifact hashes", "fitted model-artifact hashes"),
    (
        "Raw public records remain outside Git history and are handled under their documented licenses. Published repository evidence contains aggregate counts, hashes, protocols, and code rather than redistributed processed row-level outputs.",
        "Licensed raw records, prepared partitions, fitted artifacts, row-level predictions and request records are distributed in the versioned research release, outside ordinary Git history. Its inventories distinguish exact files, disclosed administrative projections and hash-only private execution files. Dataset URLs remain inert strings; publication does not authorize visiting phishing infrastructure.",
    ),
]

APPENDIX = [
    (
        "This appendix is a reader's guide",
        "The accompanying repository connects the manuscript to aggregate results, analysis code and the licensed retained-record release cited as Tallam (2026). Use release tag research-record-2026-10-04 for the exact data and measured-code record. Paths in Table A.1 are relative to dissertation/2026-10-04; unqualified data filenames refer to its aggregate-data directory. The current EVIDENCE_MAP.md maps all 27 tables and four figures to their sources and the 25-slide presentation to the same evidence. Literature citations support external claims; repository locators identify this study's artifacts, not independent corroboration.",
    ),
    (
        "The measured academic code revision is",
        "The original measured code revision is 77d128377ce5b401437d7179f5cd78fb4294b72c; the bounded follow-up revision is ef8ba5f0b357cf3dd60c4d663e6297d13334460c. The repository's research-archive/2026-10-04 guide supplies download, checksum, materialization and saved-observation recomputation commands. Its catalog inventories licensed raw data, prepared splits, fitted artifacts, predictions, requests, freezes and interrupted attempts. Private execution capabilities and full host/process logs remain excluded or explicitly projected. Checksums establish identity, and recomputation checks arithmetic; neither constitutes new model fitting, timing or independent empirical replication.",
    ),
    (
        "The bounded follow-up has a distinct source identity",
        "Public method records are in provenance/followup, with detection results in aggregate-data/detection-D and complete service results in aggregate-data/service-S. Materialized follow-up inputs, manifests, model freezes and population admissions are under dissertation/followup-20261001 in the archive. The detection supplement retains two model rows, twenty calibration bins, two invariance rows and domain sizes. Its separate arithmetic implementation checks counts, rank-based AUC, average precision, Brier score and domain-bootstrap summaries. Candidate fitting used Python 3.10.19, NumPy 2.2.6, SciPy 1.15.3 and scikit-learn 1.7.2. Historical analysis-scripts retain controlled-workspace dependencies; the public recomputation adapter and data dictionaries specify the supported use.",
    ),
    (
        "The subsequently completed broad regression run records",
        "Software checks have a separate chronology. The follow-up freeze recorded 215 focused passing tests while a broader run was in progress. That run finished with 13,016 passed, three failed and two skipped: a mock client lacked the async aclose method, preventing the intended SIGINT injection. A separately retained fixture-only correction passed all six original interruption cases, including three previously passing shift cases. The frozen source, original failure log and initial configuration-selection error before collection remain preserved. Current GitHub Actions results identify their tested commits; they do not relabel historical failures or establish manuscript acceptance. No advisor or institutional approval is inferred from these checks.",
    ),
]

TABLES = [
    (
        "2.1",
        "Closest literature and contribution boundaries",
        [
            "manuscript/Tallam_Krti_Praxis_Integrated_2026-10-04.md",
            "literature/README.md",
            "literature/source-claims.json",
        ],
        "Literature synthesis; author–date citations in each row and the manuscript bibliography remain authoritative.",
    ),
    (
        "3.1",
        "Component-to-experiment map",
        [
            "provenance/CODE_AND_ENVIRONMENT.json",
            "provenance/original/measured-source/data/rq1-transformer-cascade-contract-v2.json",
            "provenance/original/measured-source/data/rq2-gmm-development-contract-v1.json",
        ],
        "Design map, not an empirical result.",
    ),
    (
        "4.1",
        "External composition",
        ["aggregate-data/source-contingency.csv"],
        "Source/tier strata and reference-label roles.",
    ),
    (
        "4.2a",
        "Primary confusion counts",
        ["aggregate-data/primary-results.json", "aggregate-data/primary-gates.csv"],
        "Internal and external strata use distinct denominators.",
    ),
    (
        "4.2b",
        "Primary rates and bounds",
        ["aggregate-data/primary-results.json", "aggregate-data/primary-gates.csv"],
        "Observed FPR and one-sided exact bound are distinct.",
    ),
    (
        "4.3",
        "Six paired recall contrasts",
        ["aggregate-data/paired-contrasts.csv"],
        "Domain-clustered intervals; contrasts are percentage points.",
    ),
    (
        "4.4",
        "GMM training selection",
        [
            "aggregate-data/complete-secondary-results.json",
            "provenance/original/measured-source/data/rq2-gmm-development-contract-v1.json",
        ],
        "historical.rq2-gmm-development-v1-summary.json.data.candidates contains the six retained training fits.",
    ),
    (
        "4.5",
        "Monitor audit and external alerts",
        [
            "aggregate-data/complete-secondary-results.json",
            "aggregate-data/external-monitor-windows.csv",
        ],
        "Overlapping window rates, not URL-level false-positive rates.",
    ),
    (
        "4.6",
        "Twenty-five operational groups",
        [
            "aggregate-data/operational-groups.csv",
            "aggregate-data/operational-runs.csv",
        ],
        "Five-repeat pooled quantiles.",
    ),
    (
        "4.7",
        "Throughput and drain ranges",
        ["aggregate-data/operational-runs.csv"],
        "Ranges across repeats, not uncertainty intervals.",
    ),
    (
        "4.8",
        "Ranking and calibration",
        ["aggregate-data/secondary-metrics.csv", "aggregate-data/calibration-bins.csv"],
        "Selected rows; complete metrics remain in CSV.",
    ),
    (
        "4.9",
        "Additional tiers and controls",
        ["aggregate-data/secondary-metrics.csv"],
        "Tranco alerts are label-free.",
    ),
    (
        "4.10",
        "All permutation comparators",
        [
            "aggregate-data/secondary-metrics.csv",
            "aggregate-data/complete-secondary-results.json",
        ],
        "Consumed-label provenance limitation remains; no seed selection.",
    ),
    (
        "4.11",
        "Transformer seed sensitivity",
        [
            "aggregate-data/secondary-metrics.csv",
            "aggregate-data/seed-logical-invocations.csv",
        ],
        "Accepted secondary operating points, not primary replacements.",
    ),
    (
        "4.12",
        "McNemar and Holm contrasts",
        ["aggregate-data/primary-results.json"],
        "Positive-only contrasts; numerical underflow is not exact zero.",
    ),
    (
        "4.13",
        "Development probes",
        [
            "aggregate-data/complete-secondary-results.json",
            "aggregate-data/probe-decisions-and-scores.csv",
            "aggregate-data/probe-monitors-and-scores.csv",
        ],
        "Label-free development accounting.",
    ),
    (
        "4.14",
        "All twenty-two primary gates",
        ["aggregate-data/primary-gates.csv"],
        "Nine pass, thirteen fail; unrounded operands govern.",
    ),
    (
        "4.15",
        "D population admission",
        [
            "aggregate-data/detection-D/verification.json",
            "aggregate-data/detection-D/domain-size-distribution.csv",
        ],
        "Exclusion-reason incidences overlap.",
    ),
    (
        "4.16",
        "D frozen-threshold comparison",
        [
            "aggregate-data/detection-D/detection-metrics.csv",
            "aggregate-data/detection-D/verification.json",
        ],
        "8,622 eligible records; recall/specificity tradeoff.",
    ),
    (
        "4.17",
        "D requirements",
        [
            "aggregate-data/detection-D/verification.json",
            "aggregate-data/detection-D/scheme-invariance.csv",
        ],
        "Exact invariance does not change the unsupported D conjunction.",
    ),
    (
        "4.18",
        "Ten primary S pairs",
        ["aggregate-data/service-S/primary-pairs.csv"],
        "Success-only p95; paired ratio, not pooled-quantile ratio.",
    ),
    (
        "4.19",
        "Eight S groups",
        ["aggregate-data/service-S/group-metrics.csv"],
        "Synthetic service workload; unchanged structural scorer.",
    ),
    (
        "4.20",
        "Five S requirements",
        [
            "aggregate-data/service-S/requirements.csv",
            "aggregate-data/service-S/verification.json",
        ],
        "Four pass; strict response agreement fails.",
    ),
    (
        "5.1",
        "Technical contributions",
        [
            "aggregate-data/primary-gates.csv",
            "aggregate-data/detection-D/verification.json",
            "aggregate-data/service-S/verification.json",
        ],
        "Interpretive synthesis of the cited measurements.",
    ),
    (
        "5.2",
        "Decisions at each claim level",
        [
            "aggregate-data/primary-gates.csv",
            "aggregate-data/detection-D/verification.json",
            "aggregate-data/service-S/requirements.csv",
        ],
        "Original H1–H3 remain distinct from D and S.",
    ),
    (
        "A.1",
        "Evidence reading map",
        [
            "provenance/advisor-deck-scope-crosswalk-20261001.md",
            "aggregate-data/primary-gates.csv",
        ],
        "Navigation index, not additional data.",
    ),
    (
        "B.1",
        "Full eighty-arm S schedule",
        ["aggregate-data/service-S/arm-metrics.csv"],
        "The interrupted 45-arm schedule is never pooled.",
    ),
]
FIGURES = [
    (
        "3.1",
        "Implemented dataflow",
        [
            "manuscript/gwu-system-dataflow-20261001.pdf",
            "analysis-scripts/dissertation/draw_system_dataflow_20261001.py",
        ],
        "Schematic of frozen stages and accounting boundaries, not measured performance.",
    ),
    (
        "3.2",
        "Source and partition roles",
        [
            "manuscript/followup-source-partitions.pdf",
            "provenance/followup/comparison-specification-v1.md",
            "aggregate-data/detection-D/verification.json",
        ],
        "Permitted information flow; overlap screening is not temporal validation.",
    ),
    (
        "4.1",
        "Paired D recall difference",
        [
            "manuscript/followup-paired-recall.pdf",
            "aggregate-data/detection-D/verification.json",
        ],
        "paired_uncertainty.recall_difference; effect and 97.5% interval in percentage points.",
    ),
    (
        "4.2",
        "D calibration and bin populations",
        [
            "manuscript/followup-calibration.pdf",
            "aggregate-data/detection-D/calibration-bins.csv",
        ],
        "Fixed bins and all counts; empty bins have no reliability point.",
    ),
]
SLIDE_SOURCES = [
    ["RESEARCH_STORY.md"],
    [
        "provenance/advisor-deck-scope-crosswalk-20261001.md",
        "aggregate-data/primary-gates.csv",
    ],
    ["aggregate-data/source-contingency.csv", "aggregate-data/verification.json"],
    ["aggregate-data/primary-results.json"],
    ["aggregate-data/paired-contrasts.csv"],
    ["aggregate-data/primary-gates.csv", "aggregate-data/external-monitor-windows.csv"],
    ["aggregate-data/primary-gates.csv", "aggregate-data/operational-runs.csv"],
    ["aggregate-data/operational-groups.csv"],
    [
        "aggregate-data/secondary-verification.json",
        "aggregate-data/complete-secondary-results.json",
    ],
    ["provenance/followup/comparison-specification-v1.md"],
    ["provenance/followup/diagnosis-and-design.md"],
    ["aggregate-data/detection-D/verification.json"],
    [
        "aggregate-data/detection-D/detection-metrics.csv",
        "aggregate-data/detection-D/verification.json",
    ],
    ["aggregate-data/detection-D/scheme-invariance.csv"],
    [
        "aggregate-data/service-S/verification.json",
        "aggregate-data/service-S/requirements.csv",
    ],
    ["aggregate-data/service-S/primary-pairs.csv"],
    ["aggregate-data/service-S/group-metrics.csv"],
    ["RESEARCH_STORY.md"],
    ["VERIFICATION.md", "provenance/followup/service-recovery-amendment-v2.md"],
    ["aggregate-data/primary-gates.csv", "aggregate-data/paired-contrasts.csv"],
    ["aggregate-data/primary-gates.csv"],
    ["aggregate-data/primary-gates.csv"],
    ["aggregate-data/operational-groups.csv", "aggregate-data/operational-runs.csv"],
    ["aggregate-data/operational-groups.csv", "aggregate-data/operational-runs.csv"],
    ["aggregate-data/operational-groups.csv", "aggregate-data/operational-runs.csv"],
]


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def replace_span(paragraph, old, new):
    assert paragraph.text.count(old) == 1
    start = paragraph.text.index(old)
    end = start + len(old)
    offset = 0
    inserted = False
    for run in paragraph.runs:
        original = run.text
        run_end = offset + len(original)
        if offset < end and run_end > start:
            run.text = (
                original[: max(0, start - offset)]
                + (new if not inserted else "")
                + original[max(0, end - offset) :]
            )
            inserted = True
        offset = run_end
    assert inserted


def build(repository, output):
    output.mkdir(parents=True, exist_ok=False)
    source = output / "source"
    source.mkdir()
    package = repository / PACKAGE_PATH
    names = [
        "manuscript/Tallam_Krti_Praxis_Public_2026-10-04.docx",
        "manuscript/Tallam_Krti_Praxis_Citations_Expanded_2026-10-04.md",
        "manuscript/Tallam_Krti_Praxis_Citations_Expanded_2026-10-04.pdf",
        "advisor/Tallam_Praxis_Advisor_Complete_2026-10-02.pptx",
        "advisor/Tallam_Praxis_Advisor_Complete_2026-10-02.pdf",
    ]
    for name in names:
        (source / Path(name).name).write_bytes(
            subprocess.check_output(
                [
                    "git",
                    "-C",
                    str(repository),
                    "show",
                    f"{BASE_COMMIT}:{PACKAGE_PATH}/{name}",
                ]
            )
        )
    manuscript = output / "manuscript"
    manuscript.mkdir()
    for path in (package / "manuscript").iterdir():
        if not path.name.startswith("Tallam"):
            (manuscript / path.name).write_bytes(path.read_bytes())
    document = Document(source / Path(names[0]).name)
    markdown = (source / Path(names[1]).name).read_text()
    edits = list(EDITS)
    for prefix, new in APPENDIX:
        matches = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.startswith(prefix)
        ]
        assert len(matches) == 1
        edits.append((matches[0], new))
    edits.append(("Sections 1.3–1.4", "Sections 1.4–1.5"))
    paragraphs = document.paragraphs + [
        paragraph
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    ]
    for old, new in edits:
        matches = [paragraph for paragraph in paragraphs if old in paragraph.text]
        assert len(matches) == 1 and markdown.count(old) == 1, old
        replace_span(matches[0], old, new)
        markdown = markdown.replace(old, new, 1)
    anchor = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("Tibshirani, R.")
    )
    element = deepcopy(anchor._p)
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.clear()
    paragraph.add_run("Tallam, K. (2026). ")
    paragraph.add_run(TITLE).italic = True
    paragraph.add_run(
        f" (research-record-2026-10-04) [Computer software and data set]. GitHub. {RELEASE}"
    )
    assert paragraph.text == REFERENCE
    markdown = markdown.replace(
        "\n\nTibshirani, R.", "\n\n" + REFERENCE + "\n\nTibshirani, R.", 1
    )
    target = manuscript / "Tallam_Krti_Praxis_Integrated_2026-10-04.docx"
    document.save(target)
    target.with_suffix(".md").write_text(markdown)
    ledger = {
        "base_commit": BASE_COMMIT,
        "data_release": RELEASE,
        "source_files": {path.name: digest(path) for path in source.iterdir()},
        "edits": [{"old": old, "new": new} for old, new in edits],
        "added_reference": REFERENCE,
        "reference_before": "Tibshirani, R.",
        "scope": "Editorial repository integration; no scientific measurement or decision changed.",
    }
    (output / "editorial-ledger.json").write_text(
        json.dumps(ledger, indent=2, ensure_ascii=False) + "\n"
    )
    notes = integrate_deck(source / Path(names[3]).name, output)
    advisor = output / "advisor"
    (advisor / "Tallam_Praxis_Advisor_Integrated_2026-10-04.pdf").write_bytes(
        (source / Path(names[4]).name).read_bytes()
    )
    mapping = {
        "data_release": RELEASE,
        "base_commit": BASE_COMMIT,
        "tables": [
            dict(zip(("id", "title", "sources", "interpretation"), row))
            for row in TABLES
        ],
        "figures": [
            dict(zip(("id", "title", "sources", "interpretation"), row))
            for row in FIGURES
        ],
        "slides": notes,
    }
    (output / "evidence-map.json").write_text(
        json.dumps(mapping, indent=2, ensure_ascii=False) + "\n"
    )
    print(target)


def integrate_deck(source, output):
    advisor = output / "advisor"
    advisor.mkdir(exist_ok=True)
    replacements = {
        "followup-20261001/verified-detection-v1/": "aggregate-data/detection-D/",
        "verified-service-v2/": "aggregate-data/service-S/",
        "final-evidence-20261001/": "aggregate-data/",
        "comparison-specification-v1.md": "provenance/followup/comparison-specification-v1.md",
        "execution-manifest-v1.json": "materialized archive: dissertation/followup-20261001/execution-manifest-v1.json",
    }
    slides = []
    exported = []
    changed = {}
    with ZipFile(source) as archive:
        relationships = {
            entry.get("Id"): entry.get("Target")
            for entry in etree.fromstring(
                archive.read("ppt/_rels/presentation.xml.rels")
            )
        }
        order = etree.fromstring(archive.read("ppt/presentation.xml")).xpath(
            ".//p:sldId", namespaces=NS
        )
        for number, slide in enumerate(order, 1):
            path = "ppt/" + relationships[slide.get("{" + NS["r"] + "}id")]
            title = etree.fromstring(archive.read(path)).xpath(
                ".//a:t/text()", namespaces=NS
            )[0]
            relations_path = posixpath.join(
                posixpath.dirname(path), "_rels", posixpath.basename(path) + ".rels"
            )
            relationship = next(
                entry
                for entry in etree.fromstring(archive.read(relations_path))
                if entry.get("Type").endswith("/notesSlide")
            )
            note_path = posixpath.normpath(
                posixpath.join(posixpath.dirname(path), relationship.get("Target"))
            )
            note = etree.fromstring(archive.read(note_path))
            body = next(
                shape
                for shape in note.xpath(".//p:sp", namespaces=NS)
                if shape.xpath(".//p:ph[@type='body']", namespaces=NS)
            )
            frame = body.find("p:txBody", NS)
            texts = [
                "".join(paragraph.xpath(".//a:t/text()", namespaces=NS))
                for paragraph in frame.findall("a:p", NS)
            ]
            for old, new in replacements.items():
                texts = [text.replace(old, new) for text in texts]
            texts += [
                "Public sources: " + "; ".join(SLIDE_SOURCES[number - 1]),
                "Public evidence base: " + TAG_BASE + PACKAGE_PATH + "/",
                "Retained data, models and freezes: " + RELEASE,
            ]
            for paragraph in frame.findall("a:p", NS):
                frame.remove(paragraph)
            for text in texts:
                paragraph = etree.SubElement(frame, "{" + NS["a"] + "}p")
                run = etree.SubElement(paragraph, "{" + NS["a"] + "}r")
                etree.SubElement(run, "{" + NS["a"] + "}t").text = text
            changed[note_path] = etree.tostring(
                note, xml_declaration=True, encoding="UTF-8", standalone=True
            )
            exported.append(f"SLIDE {number}\n" + "\n".join(texts))
            slides.append(
                {
                    "number": number,
                    "title": title,
                    "slide_part": path,
                    "note_part": note_path,
                    "sources": SLIDE_SOURCES[number - 1],
                }
            )
        destination = advisor / "Tallam_Praxis_Advisor_Integrated_2026-10-04.pptx"
        with ZipFile(destination, "w") as revised:
            for item in archive.infolist():
                revised.writestr(
                    item, changed.get(item.filename, archive.read(item.filename))
                )
    (
        advisor / "Tallam_Praxis_Advisor_Integrated_2026-10-04_Speaker_Notes.txt"
    ).write_text("\n\n".join(exported) + "\n")
    return slides


if __name__ == "__main__":
    build(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())
