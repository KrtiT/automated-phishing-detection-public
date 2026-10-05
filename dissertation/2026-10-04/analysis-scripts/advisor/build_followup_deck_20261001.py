"""Integrate verified engineering iteration into a separate advisor working deck."""

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

import build_advisor_results_20261001 as base

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "dissertation/followup-20261001"
ORIGINAL = ROOT / "deliverables/Tallam_Praxis_Advisor_Results_2026-10-01.pptx"
OUTPUT = EVIDENCE / "advisor-work/Tallam_Praxis_Engineering_Working_2026-10-01.pptx"


def annotate(slide, explanation):
    base.LAYOUT.add_speaker_notes(slide, [
        explanation,
        "Evidence: followup-20261001/verified-detection-v1/verification.json; comparison-specification-v1.md; execution-manifest-v1.json. Original findings: final-evidence-20261001/.",
        "The bounded extension was specified October 1, 2026 after the initial results and before retrieval, fitting and scoring of its additional benchmark. It is development-informed engineering iteration, not an original confirmatory hypothesis. Operator authorization is not advisor or institutional approval. Original H1–H3 decisions are unchanged. S is not yet measured; no diagnostic timing is final evidence.",
    ])


def compose():
    evidence = json.loads((EVIDENCE / "verified-detection-v1/verification.json").read_text())
    if evidence["status"] != "verified" or evidence["records"] != 8622 or evidence["D_requirement_met"] is not False:
        raise ValueError("Unexpected or unverified follow-up evidence")
    for model, expected in [("baseline", (4619, 3940, 31, 32)), ("candidate", (4435, 3689, 282, 216))]:
        counts = evidence["metrics"][model]["counts"]
        if tuple(counts[key] for key in ("tp", "fp", "tn", "fn")) != expected:
            raise ValueError("Deck narrative does not match verified counts")
    presentation = Presentation(ORIGINAL)
    if len(presentation.slides) != 17:
        raise ValueError("Unexpected preserved deck structure")
    slide = presentation.slides[0]
    base.LAYOUT.clear_slide_body(slide)
    for shape in slide.shapes:
        if shape.has_text_frame:
            shape.text_frame.clear()
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.4), Inches(1.65), Inches(5.45), Inches(3.9))
    panel.fill.solid()
    panel.fill.fore_color.rgb = base.LAYOUT.NAVY
    panel.line.fill.background()
    base.add_text(slide, .65, 1.85, 5, 1.15, "Automated Phishing Detection\nfor Frontier AI Inference", size=24, bold=True, color=base.LAYOUT.WHITE)
    base.add_text(slide, .65, 3.15, 5, .5, "Krti Tallam | October 1, 2026", size=16, color=base.LAYOUT.WHITE)
    base.add_text(slide, .65, 3.85, 5, 1.4, "Design → evaluation → diagnosis\nModification → measured comparison\nEngineering working revision", size=17, color=base.LAYOUT.WHITE)
    annotate(slide, "Lead with the engineering investigation and its supported contributions. The original study is complete. The bounded detection extension is measured and separately recomputed. The controlled service extension is not yet measured, so this revision is explicitly working rather than final.")

    for index in (9, 10):
        base.LAYOUT.clear_slide_body(presentation.slides[index])
    base.panels(presentation.slides[9], "Technical value: isolate the mechanism and test the change",
                "Demonstrated contributions", ["Structural features: +64.07 pp internal recall over length alone.", "Complete prediction → routing → HTTP evidence on one artifact.", "Scheme-neutral representation: exact paired invariance."],
                "Engineering decisions enabled", ["Do not equate a shift alert with useful remediation.", "Require external risk evidence alongside input stability.", "Separate model cost from client and transport cost."],
                "Value: implemented contracts, controlled comparisons and reproducible design knowledge; no first-ever algorithm claim.")
    annotate(presentation.slides[9], "Initial evidence includes all 125 cells, 25 operational groups and 22 primary checks; nine component checks pass. The follow-up adds an implemented scheme-neutral feature contract and quantifies its transfer tradeoff. Scheme normalization, logistic regression and connection reuse have prior art. Contribution is the connected artifact and evidence, not invention of those methods or proof of deployment readiness.")
    base.panels(presentation.slides[10], "Current evidence and the remaining completion step",
                "Verified and integrated", ["Original RQ1–RQ3 answers and H1–H3 decisions preserved.", "8,622-row comparison; complete metrics and clustered intervals.", "Updated GWU working manuscript, figures, references and deck."],
                "To finish this revision", ["Run the fixed service comparison once under required conditions.", "Verify every scheduled arm and retain all adverse outcomes.", "Reconcile final service findings and complete document QA."],
                "S: not yet measured. No final extension or degree-acceptance claim; the completed original package is preserved.")
    annotate(presentation.slides[10], "The only new research measurement still outstanding is the fixed service comparison. Timing requires actual AC power and no competing computational workload. The manuscript keeps the original credentials and original files; this deck and its new DOCX/PDF are distinct working copies until the extension and artifact verification finish.")

    added = [presentation.slides.add_slide(presentation.slides[1].slide_layout) for _ in range(6)]
    for slide in added:
        base.LAYOUT.clear_slide_body(slide)
    base.panels(added[0], "One engineering process, with distinct evidence stages",
                "Design and evaluate", ["Define URL, risk, routing and service contracts.", "Measure all 22 primary checks and 25 operational groups.", "Identify where component benefits meet or miss joint requirements."],
                "Diagnose and compare", ["Trace scheme spelling through structural features.", "Implement a transport-neutral representation; freeze its threshold.", "Measure an additional benchmark; isolate HTTP client topology."],
                "Original H1–H3 remain unchanged. The bounded extension adds measured mechanism-level evidence, not rewritten history.")
    annotate(added[0], "This diagram organizes the work as engineering iteration. Original H1/H2/H3 remain not supported under unchanged conjunctive gates. Nine original component checks pass. The D requirement is separate, and its exact-invariance diagnostic is not retroactively promoted to a new confirmatory hypothesis. Methods record the actual chronology once.")

    base.panels(added[1], "Diagnosis selects two bounded interventions",
                "Representation pathway", ["Scheme affects is_https, length, ratios and entropy.", "Dropping is_https alone does not remove scheme information.", "Use one modeling prefix for every feature; omit is_https."],
                "Service pathway", ["No-model diagnostics expose client/transport overhead.", "Compare shared versus worker-owned persistent clients.", "Keep service, scorer, inputs and request limits unchanged."],
                "Diagnostics motivate tests; they do not establish a single cause. Transformer, GMM and routing are not retuned.")
    annotate(added[1], "Original descriptive is_https PSI motivates investigation but does not prove the cause of external false positives. Transport-neutral inputs preserve raw URLs, canonicalization and overlap rules. Shared versus worker-owned clients change client topology only. Instrumented development timings are excluded from final S evidence. Technical grounding includes Arp et al. (2022), Pendlebury et al. (2019), Guo et al. (2017) and Geifman and El-Yaniv (2019); these references do not establish novelty for the intervention.")

    base.panels(added[2], "Detection comparison: freeze first, evaluate every admitted row",
                "Additional benchmark", ["Publisher 2020 benchmark: 11,430 source rows.", "8,622 admitted: P=4,651; N=3,971; 6,273 domains.", "2,808 exclusions; overlap and quarantine records retained."],
                "Fixed comparison", ["Unchanged 25-feature L1 comparator and threshold.", "24-feature candidate: one fit; development-only calibration.", "All admitted rows scored once; paired domain bootstrap."],
                "Retrospective, publisher-labeled comparison; shared upstream feeds. Domain exclusion is not temporal independence.")
    annotate(added[2], "Mendeley DOI 10.17632/c2gw7fy2j4.3, CC BY 4.0. Publisher metadata says May 2020 while the associated preprint says March; use '2020 benchmark.' Quarantine reason incidences overlap and must not be summed as distinct excluded rows. Largest cluster: 351 records. Candidate fit: 166,248 training rows, convergence at 4,971/5,000 iterations. Validation: 32,695 rows; one-sided CP FPR upper bound 0.9968%. Candidate threshold 0.6815549262616749; comparator 0.2670846328466124. Both thresholds froze before additional-benchmark predictions.")

    slide = added[3]
    base.LAYOUT.set_title(slide, "D: full operating-point tradeoff at the frozen thresholds")
    metrics = evidence["metrics"]
    rows = []
    for label, key in [("Recall", "recall"), ("False-positive rate", "fpr"), ("Precision", "precision"), ("ROC AUC", "roc_auc"), ("Average precision", "average_precision"), ("Brier score", "brier")]:
        values = [metrics[model][key] for model in ("baseline", "candidate")]
        formatted = [f"{100 * value:.2f}%" if key in {"recall", "fpr", "precision"} else f"{value:.4f}" for value in values]
        rows.append([label, *formatted])
    base.grid(slide, ["Metric", "Unchanged L1", "Scheme-neutral"], rows, widths=[.42, .29, .29], size=13, top=1.25, height=3.05)
    base.add_text(slide, .55, 4.43, 8.8, .7, "251 fewer false positives; 184 fewer detected positives.\nD: not supported — recall gain and the 1% FPR requirement are not met.", size=15, bold=True)
    base.LAYOUT.add_note(slide, "Recall difference −3.96 pp; 97.5% domain-cluster interval [−5.05, −3.00] pp. No threshold retuning.", y=5.32, font_size=11)
    annotate(slide, "Comparator counts: TP 4,619; FN 32; FP 3,940; TN 31. Candidate: TP 4,435; FN 216; FP 3,689; TN 282. Candidate FPR 97.5% interval [91.8651%, 93.9100%]. D requires at least five points of recall gain, a lower 97.5% interval bound above zero and observed candidate FPR at most 1%; all three fail. The candidate satisfies the representation check separately. Ten thousand paired registrable-domain resamples; seed 20261001; zero undefined replicates. Ranking and Brier changes are descriptive, not additional significance claims. Separate arithmetic recomputation verifies every summary; it is not external independent replication.")

    base.panels(added[4], "Achieved property: exact scheme invariance",
                "Prescribed metamorphic check", ["Swap HTTP ↔ HTTPS; keep the remainder unchanged.", "Candidate features and scores match on 8,622 / 8,622 pairs.", "Candidate decisions: zero flips."],
                "What the comparison establishes", ["Unchanged L1: 481 decision flips; 2,909 exact score matches.", "The modified feature path removes the intended dependency.", "Stable representation alone does not ensure low-FPR transfer."],
                "Companions test representation, not additional labeled websites. Comparator flips are not automatically classification errors.")
    annotate(added[4], "This is a positive, exact implementation-conformance result under a specified transformation. Removing scheme from every affected feature, not merely its binary indicator, gives the property. Maximum comparator absolute score difference: 0.99664935. All raw observations remain intact. This result must not be described as successful detector repair, prospective temporal validation or a newly passing original hypothesis.")

    base.panels(added[5], "S: controlled client-topology comparison is the remaining test",
                "Fixed 80-arm design", ["No-model control and unchanged structural scorer.", "c1 and c64; ten pairs; alternating shared/worker order.", "Fresh services; 1,000 warmups + 10,000 requests per arm."],
                "Primary structural c64 requirement", ["Median p95 ratio ≤0.8; 97.5% upper interval <1.", "Worker pooled p95 ≤200 ms; errors <0.1%.", "Exact response agreement apart from request IDs."],
                "S: not yet measured. Actual AC and a quiet reserved host are required; diagnostic timings do not qualify.")
    annotate(added[5], "Ten paired run-bootstrap resamples are NOT the design: the design uses ten operational pairs and 10,000 bootstrap replicates, seed 20261002. Full response agreement includes admission sequence; prediction-only agreement is separately reported. Input strings use reserved .example hosts, carry no labels or prevalence claim, and are not fetched. No transformer is invoked. An environmental interruption stops the schedule and preserves incomplete evidence rather than authorizing a retry. This comparison cannot replace original H3 measurements.")

    slide_ids = presentation.slides._sldIdLst
    original_ids = list(slide_ids)[:17]
    added_ids = list(slide_ids)[17:]
    for identifier in list(slide_ids):
        slide_ids.remove(identifier)
    for identifier in original_ids[:9] + added_ids + original_ids[9:]:
        slide_ids.append(identifier)
    for slide in list(presentation.slides)[1:]:
        title = slide.shapes.title
        title.left, title.top = Inches(.48), Inches(.22)
        title.width, title.height = Inches(9.02), Inches(.85)
    properties = presentation.core_properties
    properties.title = "Praxis Engineering Working Revision — October 1, 2026"
    properties.subject = "Original RQ1–RQ3; verified detection extension; service comparison pending"
    properties.author = properties.last_modified_by = "Krti Tallam"
    properties.modified = datetime.now(timezone.utc)
    return presentation


def main():
    original_hash = hashlib.sha256(ORIGINAL.read_bytes()).hexdigest()
    presentation = compose()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    if hashlib.sha256(ORIGINAL.read_bytes()).hexdigest() != original_hash:
        raise ValueError("Original deck changed")
    receipt = {
        "status": "working_deck_built_service_pending",
        "slides": len(presentation.slides),
        "preserved_original_sha256": original_hash,
        "pptx_sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
        "detection_verification_sha256": hashlib.sha256((EVIDENCE / "verified-detection-v1/verification.json").read_bytes()).hexdigest(),
        "render_review": "not_yet_complete",
    }
    (OUTPUT.parent / "build-receipt.json").write_text(json.dumps(receipt, indent=2) + "\n")
    print(json.dumps(receipt, indent=2))


if __name__ == "__main__":
    main()
