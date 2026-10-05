import importlib
import importlib.util
import unittest

from pptx import Presentation


def slide_text(slide):
    values = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            values.append(shape.text)
        if shape.has_table:
            values.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(values)


class FollowupDeckTests(unittest.TestCase):
    def setUp(self):
        module_name = "build_followup_deck_20261001"
        self.assertIsNotNone(importlib.util.find_spec(module_name), "Follow-up builder is not implemented")
        self.builder = importlib.import_module(module_name)

    def test_preserves_original_questions_and_complete_appendices(self):
        original = Presentation(self.builder.ORIGINAL)
        current = self.builder.compose()
        self.assertEqual(len(current.slides), 23)
        self.assertEqual(slide_text(current.slides[1]), slide_text(original.slides[1]))
        for original_index, current_index in zip(range(11, 17), range(17, 23)):
            self.assertEqual(slide_text(current.slides[current_index]), slide_text(original.slides[original_index]))

    def test_reports_measured_tradeoff_and_unmeasured_service(self):
        current = self.builder.compose()
        text = "\n".join(slide_text(slide) for slide in current.slides)
        for expected in ["8,622", "251", "184", "92.90%", "95.36%", "481", "97.5%", "D: not supported", "S: not yet measured"]:
            self.assertIn(expected, text)
        self.assertIn("H1–H3", text)
        self.assertNotIn("Finished local deliverables", text)

    def test_chronology_and_provenance_in_notes(self):
        current = self.builder.compose()
        notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in current.slides)
        for expected in ["after the initial results", "verified-detection-v1/verification.json", "comparison-specification-v1.md", "not advisor"]:
            self.assertIn(expected, notes)
        self.assertNotEqual(self.builder.OUTPUT, self.builder.ORIGINAL)
        self.assertIn("Working", self.builder.OUTPUT.name)


if __name__ == "__main__":
    unittest.main()
