"""Verify current render identities and navigation; export deck notes and QA views."""

import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

import fitz
from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation

HERE = Path(__file__).resolve().parent
CONTEXT = HERE.parent
REVIEW = CONTEXT / "reviews/gwu-native-render-20261001"
MANUSCRIPT = HERE / "Tallam_Krti_Praxis_GWU_2026-10-01.docx"
DECK = CONTEXT / "deliverables/Tallam_Praxis_Advisor_Results_2026-10-01.pptx"
WORD = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
PRESERVED = {
    "attachments/pduS8C/DEng Praxis Template - Online Programs Rev 2026.docx":
        "0c22a6ba142187a3b763f32045bf86a59a4e747dbb32b6335c53a3e3033c2547",
    "attachments/foUeB5/Tallam_Krti_Praxis_2026-06-15.docx":
        "96c056e6becf2bab7335adf4ed850707ca049749233b819fde7fb80a060daf3d",
    "deliverables/Tallam_Praxis_Advisor_Update_2026-09-17_Meeting_Final.pptx":
        "6b6fc2a36b9cbf50061ca1865c449e87c3492f7733394929f8e32b4934cdf5de",
}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    for name, expected in PRESERVED.items():
        if digest(CONTEXT / name) != expected:
            raise ValueError(f"Preserved source changed: {name}")
    pages = json.loads((HERE / "gwu-pagination-map-20261001.json").read_text())
    entries = json.loads((HERE / "gwu-navigation-20261001.json").read_text())
    with ZipFile(MANUSCRIPT) as archive:
        document = etree.fromstring(archive.read("word/document.xml"))
    paragraphs = document.xpath("//w:p", namespaces=WORD)
    for entry in entries:
        bookmark = entry["bookmark"]
        matches = []
        for paragraph in paragraphs:
            instructions = paragraph.xpath(".//w:instrText/text()", namespaces=WORD)
            anchors = paragraph.xpath(".//w:hyperlink/@w:anchor", namespaces=WORD)
            if bookmark in anchors or any(f"PAGEREF {bookmark} " in item for item in instructions):
                matches.append(paragraph)
        if len(matches) != 1:
            raise ValueError(f"Nonunique navigation field: {bookmark}")
        text = "".join(matches[0].xpath(".//w:t/text()", namespaces=WORD))
        if text != entry["text"] + pages[bookmark]:
            raise ValueError(f"Stale navigation field: {bookmark}: {text}")
    manuscript_pdf = REVIEW / MANUSCRIPT.with_suffix(".pdf").name
    render = json.loads((REVIEW / "render-check.json").read_text())
    if digest(manuscript_pdf) != render["pdf_sha256"]:
        raise ValueError("Manuscript PDF changed after pagination check")
    slides = Presentation(DECK)
    deck_pdf = REVIEW / DECK.with_suffix(".pdf").name
    rendered = fitz.open(deck_pdf)
    if len(slides.slides) != 17 or len(rendered) != 17:
        raise ValueError("Expected all 17 advisor slides")
    notes, outside = [], []
    for index, (slide, page) in enumerate(zip(slides.slides, rendered), 1):
        for word in page.get_text("words"):
            if word[0] < 0 or word[1] < 0 or word[2] > page.rect.width or word[3] > page.rect.height:
                outside.append({"slide": index, "word": word[:5]})
        if not page.get_text().strip():
            raise ValueError(f"Empty slide: {index}")
        notes.append(f"SLIDE {index}\n{slide.notes_slide.notes_text_frame.text.strip()}")
        page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(REVIEW / f"deck-{index:02}.png")
    if outside:
        raise ValueError(outside)
    for start in range(0, 17, 6):
        sheet = Image.new("RGB", (1500, 1410), "#dddddd")
        drawing = ImageDraw.Draw(sheet)
        for offset, index in enumerate(range(start, min(start + 6, 17))):
            picture = Image.open(REVIEW / f"deck-{index + 1:02}.png")
            picture.thumbnail((720, 425))
            left, top = 15 + 750 * (offset % 2), 30 + 470 * (offset // 2)
            sheet.paste(picture, (left, top))
            drawing.text((left, top - 20), f"Slide {index + 1}", fill="black")
        sheet.save(REVIEW / f"deck-contact-{start // 6 + 1:02}.png")
    notes_path = DECK.with_name(DECK.stem + "_Speaker_Notes.txt")
    notes_path.write_text("\n\n".join(notes) + "\n")
    with fitz.open(manuscript_pdf) as manuscript:
        for index in (29, 90, 91, 92):
            manuscript[index].get_pixmap(matrix=fitz.Matrix(1.7, 1.7), alpha=False).save(
                REVIEW / f"manuscript-detail-{index + 1:02}.png")
    report = {
        "status": "passed", "navigation_cached_values_checked": len(entries),
        "manuscript_pages": render["pages"], "deck_slides": 17,
        "deck_out_of_bounds_words": outside, "preserved_source_sha256": PRESERVED,
        "artifact_sha256": {path.name: digest(path) for path in
                            (MANUSCRIPT, manuscript_pdf, DECK, deck_pdf, notes_path)},
        "render_engine": "LibreOffice 26.8.0.3, local installation",
        "visual_review": "Views generated; inspection recorded separately.",
    }
    (REVIEW / "deliverables-check.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
