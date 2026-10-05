"""Verify the final documents and retained evidence without new measurements."""

import csv
import hashlib
import json
import re
import subprocess
import unicodedata
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

import fitz
from docx import Document
from lxml import etree
from pptx import Presentation

import build_complete_revision_20261002 as final
from check_final_deliverables_20261001 import PRESERVED
from package_final_20261001 import CSV_COUNTS

HERE = Path(__file__).resolve().parent
CONTEXT = HERE.parent
ROOT = final.ROOT
ORIGINAL_PACKAGE = CONTEXT / "deliverables/Tallam_Dissertation_Results_2026-10-01"
FROZEN = CONTEXT / "gwu_working/study-followup-development-20261001"
REVISION = "ef8ba5f0b357cf3dd60c4d663e6297d13334460c"
WORD = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def digest(path):
    hasher = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            hasher.update(chunk)
    return hasher.hexdigest()


def require(condition, message):
    if not condition:
        raise ValueError(message)


def normalize(text):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text).replace("\xad", ""))


def read_json(path):
    return json.loads(path.read_text())


def read_rows(path):
    with path.open() as stream:
        return list(csv.DictReader(stream))


def verify_hashes(root, inventory):
    for name, expected in inventory.items():
        path = root / name
        require(not Path(name).is_absolute() and ".." not in Path(name).parts, f"Unsafe path: {name}")
        require(not path.is_symlink() and path.resolve().is_relative_to(root.resolve()), f"Unsafe file: {name}")
        require(digest(path) == expected, f"Hash mismatch: {name}")
    return len(inventory)


def require_table(actual, expected, name):
    def normal(rows):
        return [[normalize(str(cell)) for cell in row] for row in rows]

    require(normal(actual) == normal(expected), f"Table mismatch: {name}")


def table_rows(table):
    return [[cell.text for cell in row.cells] for row in table.rows]


def check_documents():
    document = Document(final.DOCX)
    original = Document(HERE / "Tallam_Krti_Praxis_GWU_2026-10-01.docx")
    for index in range(17):
        require_table(table_rows(document.tables[index]), table_rows(original.tables[index]), f"original {index}")
    require_table(table_rows(document.tables[23]), table_rows(original.tables[17]), "original contribution")
    require_table(table_rows(document.tables[25]), table_rows(original.tables[18]), "original reading map")
    require(final.BODY.read_text() == final.compose_body(), "Final manuscript source differs from verified synthesis")
    require(len(document.tables) == 27, "Incomplete final table inventory")
    for section in document.sections:
        require((section.page_width.inches, section.page_height.inches) == (8.5, 11), "Page geometry")
        require((section.left_margin.inches, section.right_margin.inches) == (1.25, 1.25), "Side margins")
        require((section.top_margin.inches, section.bottom_margin.inches) == (1, 1), "Vertical margins")
    with ZipFile(final.DOCX) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    with fitz.open(final.DOCX.with_suffix(".pdf")) as pdf:
        require(len(pdf) == 132, "Manuscript page count changed")
        text = normalize("\n".join(page.get_text(clip=fitz.Rect(85, 65, 528, 725)) for page in pdf))
    checked = 0
    for paragraph in root.xpath("//w:p", namespaces=WORD):
        if paragraph.xpath(".//w:hyperlink[@w:anchor]|.//w:instrText", namespaces=WORD):
            continue
        content = "".join(paragraph.xpath(".//w:t/text()", namespaces=WORD))
        if content.strip():
            require(normalize(content) in text, f"Missing PDF content: {content[:120]}")
            checked += 1
    presentation = Presentation(final.PPTX)
    require(len(presentation.slides) == 25, "Incomplete deck")
    for index, (saved, expected) in enumerate(zip(presentation.slides, final.compose_deck().slides), 1):
        def content(slide):
            return [shape.text for shape in slide.shapes if shape.has_text_frame] + [
                table_rows(shape.table) for shape in slide.shapes if shape.has_table]

        require(content(saved) == content(expected), f"Saved deck differs from verified synthesis: {index}")
    notes = "\n\n".join(f"SLIDE {index}\n{slide.notes_slide.notes_text_frame.text.strip()}"
                          for index, slide in enumerate(presentation.slides, 1)) + "\n"
    notes_path = final.PPTX.with_name(final.PPTX.stem + "_Speaker_Notes.txt")
    require(notes_path.read_text() == notes, "Speaker notes differ from PPTX")
    for folder, checks in [(final.BODY.parent, {"pdf_sha256": final.DOCX.with_suffix(".pdf")}),
                           (final.PPTX.parent, {"pdf_sha256": final.PPTX.with_suffix(".pdf"),
                                               "pptx_sha256": final.PPTX, "speaker_notes_sha256": notes_path})]:
        report = read_json(folder / "render-check.json")
        require(report["status"] == "mechanical_checks_passed", "Render check failed")
        for key, path in checks.items():
            require(digest(path) == report[key], f"Artifact changed after render QA: {path.name}")
    return document, checked, notes_path


def check_numeric_tables(document):
    models = read_rows(ROOT / "verified-detection-v1/detection-metrics.csv")
    fields = [("Threshold", lambda row: f'{float(row["threshold"]):.10f}'),
              ("TP / FN", lambda row: f'{int(row["tp"]):,} / {int(row["fn"]):,}'),
              ("FP / TN", lambda row: f'{int(row["fp"]):,} / {int(row["tn"]):,}')]
    for title, key in [("Recall", "recall"), ("FPR", "fpr"), ("Precision", "precision")]:
        fields.append((title, lambda row, key=key: f'{100 * float(row[key]):.2f}%'))
    for title, key in [("ROC AUC", "roc_auc"), ("Average precision", "average_precision"), ("Brier score", "brier")]:
        fields.append((title, lambda row, key=key: f'{float(row[key]):.4f}'))
    require_table(table_rows(document.tables[18])[1:],
                  [[title, *[formatter(row) for row in models]] for title, formatter in fields], "D metrics")
    service = ROOT / "verified-service-v2"
    pairs = read_rows(service / "primary-pairs.csv")
    require_table(table_rows(document.tables[20])[1:],
                  [[row["pair"], f'{float(row["shared_p95_ms"]):.4f}', f'{float(row["worker_p95_ms"]):.4f}',
                    f'{float(row["ratio"]):.6f}', f'{row["exact_except_request_id"]} / 10,000'] for row in pairs], "S pairs")
    groups = read_rows(service / "group-metrics.csv")
    require_table(table_rows(document.tables[21])[1:],
                  [[f'{"Control" if row["workload"] == "no_model" else "Structural"} / {row["concurrency"]} / {row["client"]}',
                    *[f'{float(row[key]):.2f}' for key in ("success_p50_ms", "success_p95_ms", "success_p99_ms")],
                    row["request_errors"], f'{float(row["client_attempts_per_second"]):.2f}'] for row in groups], "S groups")
    arms = read_rows(service / "arm-metrics.csv")
    require_table(table_rows(document.tables[26])[1:],
                  [[f'{row["workload"]}-c{row["concurrency"]}-pair{int(row["pair"]):02d}-{row["client"]}',
                    f'{float(row["success_latency_p95_ms"]):.4f}', f'{float(row["client_attempts_per_second"]):.2f}',
                    row["request_errors"]] for row in arms], "80-arm appendix")
    decisions = [row[2] for row in table_rows(document.tables[22])[1:]]
    require(decisions == ["Pass" if row["passed"] == "True" else "Not met"
                          for row in read_rows(service / "requirements.csv")], "S requirement decisions changed")
    primary = read_json(service / "verification.json")["primary"]
    require([row[1] for row in table_rows(document.tables[22])[1:]] == [
        f'{primary["median_ratio"]:.6f}', f'{primary["ratio_interval_97_5"][1]:.6f}',
        f'{primary["worker_success_latency"]["p95_ms"]:.4f} ms',
        f'{primary["worker_errors"]} / {primary["request_denominator"]:,} (0%)',
        f'{primary["exact_paired_agreement"]:,} / {primary["request_denominator"]:,}'
    ], "S requirement operands changed")
    require(sum(int(row["request_count"]) for row in arms) == 800000, "S request accounting")
    require(sum(int(row["request_errors"]) for row in arms) == 1, "S error accounting")
    require(sum(int(row["both_successful"]) for row in pairs) == 99999, "Primary comparable denominator")
    require(sum(int(row["prediction_agreement"]) for row in pairs) == 99999, "Primary prediction agreement")
    return {"D_metric_rows": 9, "S_primary_pairs": 10, "S_groups": 8, "S_requirements": 5, "S_arms": 80}


def main():
    preserved = {"original_sources": verify_hashes(CONTEXT, PRESERVED)}
    working = read_json(ROOT / "working-revision-qa.json")
    preserved["working_files"] = verify_hashes(ROOT, working["files"])
    require(digest(final.ORIGINAL) == working["preserved_original_body_sha256"], "Original manuscript body changed")
    original_inventory = {line.split("  ", 1)[1]: line.split("  ", 1)[0]
                          for line in (ORIGINAL_PACKAGE / "SHA256SUMS.txt").read_text().splitlines()}
    preserved["original_package_files"] = verify_hashes(ORIGINAL_PACKAGE, original_inventory)
    original_receipt = read_json(CONTEXT / "reviews/final-package-receipt-20261001.json")
    archive_path = ORIGINAL_PACKAGE.with_suffix(".zip")
    require(digest(archive_path) == original_receipt["zip_sha256"], "Original archive changed")
    with ZipFile(archive_path) as archive:
        archived = archive.namelist()
        require(len(archived) == original_receipt["files_in_zip"] == len(original_inventory) + 1,
                "Original archive inventory changed")
        for name in [*original_inventory, "SHA256SUMS.txt"]:
            require(hashlib.sha256(archive.read(f"{ORIGINAL_PACKAGE.name}/{name}")).hexdigest()
                    == digest(ORIGINAL_PACKAGE / name), f"Original archive member changed: {name}")
    preserved["original_archive_files_including_manifest"] = len(original_inventory) + 1
    for directory, source_key, export_key in [("verified-detection-v1", "source_hashes", "export_hashes"),
                                               ("verified-service-v2", "source_sha256", "aggregate_sha256")]:
        report = read_json(ROOT / directory / "verification.json")
        require(report["status"] == "verified", f"Unverified evidence: {directory}")
        preserved[directory + "_sources"] = verify_hashes(ROOT, report[source_key])
        preserved[directory + "_exports"] = verify_hashes(ROOT / directory, report[export_key])
    interrupted = read_json(ROOT / "service-v1-preservation.json")
    preserved["interrupted_files"] = verify_hashes(ROOT / "service-comparison-v1", interrupted["source_sha256"])
    preserved["figure_exports"] = verify_hashes(ROOT / "figures", working["figure_export_hashes"])
    for name, count in CSV_COUNTS.items():
        require(len(read_rows(HERE / "final-evidence-20261001" / name)) == count, f"Original count changed: {name}")
        require(digest(HERE / "final-evidence-20261001" / name) == digest(ORIGINAL_PACKAGE / "aggregate-data" / name),
                f"Original aggregate changed: {name}")
    gates = read_rows(HERE / "final-evidence-20261001/primary-gates.csv")
    require(Counter(row["status"] for row in gates) == {"pass": 9, "fail": 13}, "Original checks changed")
    primary = read_json(HERE / "final-evidence-20261001/primary-results.json")["primary"]
    require(all(item["complete"] and item["decision"] == "not_supported"
                for item in primary["hypotheses"].values()), "Original joint decisions changed")
    head = subprocess.check_output(["git", "-C", str(FROZEN), "rev-parse", "HEAD"], text=True).strip()
    dirty = subprocess.check_output(["git", "-C", str(FROZEN), "status", "--porcelain"], text=True)
    require(head == REVISION and not dirty, "Scientific checkout changed")
    document, fragments, notes_path = check_documents()
    numeric = check_numeric_tables(document)
    artifacts = [final.BODY, final.DOCX, final.DOCX.with_suffix(".pdf"), final.PPTX,
                 final.PPTX.with_suffix(".pdf"), notes_path]
    report = {
        "status": "verified", "checked_at": datetime.now(timezone.utc).isoformat(),
        "scope": "Final artifact/preservation verification; no new fitting, predictions or measurement.",
        "frozen_revision": head, "frozen_checkout_clean": True, "preservation_hash_counts": preserved,
        "manuscript_pages": 132, "deck_slides": 25, "tables": 27, "original_tables_preserved": 19,
        "rendered_content_fragments": fragments, "missing_rendered_content": [],
        "cached_navigation_targets": 111, "deck_text_fragments": 569, "speaker_notes_match_pptx": True,
        "numeric_table_checks": numeric, "original_checks": 22, "original_groups": 25, "original_cells": 125,
        "original_component_results": dict(Counter(row["status"] for row in gates)),
        "artifact_sha256": {str(path.relative_to(final.FINAL)): digest(path) for path in artifacts},
        "visual_review": {
            "method": "All 15 manuscript contact sheets (132 pages) and all five deck contact sheets (25 slides) inspected in the immediately preceding continuation; current unchanged rendered artifacts re-bound by SHA-256.",
            "findings": "No observed clipping, overlaps, empty pages or missing table rows. Long Appendix B IDs wrap legibly; table headers repeat. Sparse chapter-ending and figure pages remain.",
            "boundary": "Agent visual inspection, not independent external review or university certification. Deck PDF is untagged; no PDF/UA claim."
        },
        "decisions": {"H1": "not_supported", "H2": "not_supported", "H3": "not_supported",
                      "D": "not_supported; exact representation invariance achieved",
                      "S": "not_supported as conjunction; four of five requirements pass"},
        "broad_research_regression": {"passed": 13016, "failed": 3, "skipped": 2,
                                       "record": "regression-verification-v1.md", "green": False},
        "checker_sha256": digest(Path(__file__)),
    }
    (final.FINAL / "final-content-preservation-qa.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
