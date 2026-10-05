"""Update the preserved advisor deck with verified final-study results."""

import csv
import hashlib
import importlib.util
import json
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "dissertation/final-evidence-20261001"
TEMPLATE = ROOT / "deliverables/Tallam_Praxis_Advisor_Update_2026-09-17_Meeting_Final.pptx"
OUTPUT = ROOT / "deliverables/Tallam_Praxis_Advisor_Results_2026-10-01.pptx"
SPEC = importlib.util.spec_from_file_location("layout", Path(__file__).with_name("build_advisor_update_2026_09_03.py"))
LAYOUT = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(LAYOUT)


def add_text(slide, x, y, width, height, text, *, size, bold=False, color=LAYOUT.TEXT):
    box = LAYOUT.add_text(slide, x, y, width, height, text, size=size, bold=bold, color=color)
    for paragraph in box.text_frame.paragraphs:
        for run in paragraph.runs:
            LAYOUT.set_run_font(run, size, bold=bold, color=color)
    return box


def records(name):
    with (EVIDENCE / name).open() as source:
        return list(csv.DictReader(source))


def notes(slide, explanation, sources):
    LAYOUT.add_speaker_notes(slide, [explanation, "Evidence: final-evidence-20261001/" + "; final-evidence-20261001/".join(sources), "All three hypotheses are complete and not supported. No failed gate, interruption or adverse result is suppressed. Operator authorization is not advisor approval."])


def panels(slide, title, left_heading, left, right_heading, right, note):
    LAYOUT.set_title(slide, title)
    LAYOUT.add_panel(slide, .48, 1.12, 4.35, 3.98, left_heading, left, font_size=14, spacing=12)
    LAYOUT.add_panel(slide, 5.17, 1.12, 4.35, 3.98, right_heading, right, font_size=14, spacing=12)
    LAYOUT.add_note(slide, note, y=5.32, height=.57, font_size=11.3)


def grid(slide, headers, rows, widths=None, size=12, top=1.3, height=3.6):
    table = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(.48), Inches(top), Inches(9.02), Inches(height)).table
    if widths:
        for column, fraction in zip(table.columns, widths):
            column.width = Inches(9.02 * fraction)
    for row_index, values in enumerate([headers, *rows]):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LAYOUT.LIGHT_BLUE if row_index == 0 else (LAYOUT.WHITE if row_index % 2 else LAYOUT.LIGHT)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(.07)
            cell.margin_top = cell.margin_bottom = Inches(.03)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.space_after = Pt(0)
                for run in paragraph.runs:
                    LAYOUT.set_run_font(run, size, bold=row_index == 0, color=LAYOUT.NAVY)


def main():
    if json.loads((EVIDENCE / "verification.json").read_text())["status"] != "verified":
        raise ValueError("Unverified primary evidence")
    original_hash = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
    presentation = Presentation(TEMPLATE)
    if len(presentation.slides) != 11:
        raise ValueError("Unexpected historical slide count")
    for slide in presentation.slides:
        LAYOUT.clear_slide_body(slide)
    while len(presentation.slides) < 17:
        presentation.slides.add_slide(presentation.slides[1].slide_layout)
    titles = []
    slide = presentation.slides[0]
    for shape in list(slide.shapes):
        if shape.has_text_frame:
            shape.text_frame.clear()
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.4), Inches(1.65), Inches(5.45), Inches(3.9))
    panel.fill.solid()
    panel.fill.fore_color.rgb = LAYOUT.NAVY
    panel.line.fill.background()
    add_text(slide, .65, 1.85, 5.0, 1.15, "Automated Phishing Detection\nfor Frontier AI Inference", size=24, bold=True, color=LAYOUT.WHITE)
    add_text(slide, .65, 3.15, 5.0, .5, "Krti Tallam | October 1, 2026", size=16, color=LAYOUT.WHITE)
    add_text(slide, .65, 3.85, 5.0, 1.4, "Completed study: 125 cells\n25 operational groups · 22 primary checks\nJoint operating gates not met", size=18, color=LAYOUT.WHITE)
    titles.append("Complete study: measured limits of the frozen system")
    notes(slide, "Lead with completion and the result, not software activity. The story connects internal versus external detection, policy routing and physical HTTP cost. H1/H2/H3 are unsupported under unchanged rules; this does not mean that every component failed. Contribution is the bounded empirical joint evaluation, not a first-ever claim.", ["verification.json", "primary-results.json"])
    slide = presentation.slides[1]
    title = "The questions promised in August and September"
    LAYOUT.set_title(slide, title)
    questions = ["RQ1: What incremental value do structural URL features and character-level representations provide under registrable-domain-disjoint and external evaluation?", "RQ2: Can GMM-based monitoring detect an external source/domain shift and guide escalation without exceeding the low-FPR operating constraint?", "RQ3: What detection, escalation, throughput, and latency tradeoffs determine whether the fixed cascade is viable inline?"]
    for index, question in enumerate(questions):
        add_text(slide, .58, 1.2 + 1.2 * index, 8.84, 1.0, question, size=17, color=LAYOUT.NAVY)
    LAYOUT.add_note(slide, "Exact active wording: August 20 and September 3 decks; carried forward September 17.", y=5.32)
    titles.append(title)
    notes(slide, "These are the active questions, not the superseded June/August 6 distillation or feature-fusion proposal. No 300M-to-30M distillation, 92% retention, +5-point AUC or 40% latency reduction is claimed. The manuscript gives direct answers using these exact questions.", ["primary-results.json"])
    slide = presentation.slides[2]
    title = "Evidence coverage and population boundaries"
    panels(slide, title, "What was measured", ["Internal: 34,593 rows; 14,326 positives and 20,267 negatives.", "External: 8,701 retained rows; gold 69, certified 2,497.", "125 operational cells; 250 owned service/client exits verified."], "What the labels mean", ["Publisher reference labels, not independent adjudication.", "Tranco: 1,163 label-free controls, never labeled FPR.", "72 retained + 53 new cells; no accepted adverse repeat replaced."], "URL-representation and two-session amendments remain disclosed; no claim of lifetime-unseen test data.")
    titles.append(title)
    notes(slide, "External source/tier details: silver 417, bronze 4,555, Tranco 1,163, gold 69 and certified 2,497; 4,150 distinct domains overall. The published test has 8,941 records, 240 excluded. Source/tier filters follow routing. Exact publisher url_norm was adopted after preparation/label-count exposure and before amended predictions; checkpoint recovery followed earlier predictions and partial runs. These timings are not advisor approval or complete pre-access prespecification.", ["verification.json", "source-contingency.csv"])
    slide = presentation.slides[3]
    title = "RQ1: internal recall does not ensure external specificity"
    LAYOUT.set_title(slide, title)
    grid(slide, ["Frozen detector", "Internal recall", "Internal FPR", "Gold recall", "Certified FPR"], [["Length-only", "34.61%", "0.5872%", "23.19%", "9.4513%"], ["Logistic-L1", "98.68%", "0.7747%", "100%", "90.5487%"], ["Fixed cascade", "98.68%", "0.7747%", "100%", "90.5487%"], ["Transformer", "99.32%", "1.0115%", "100%", "99.2791%"]], widths=[.25,.18,.18,.18,.21], size=13, height=2.85)
    add_text(slide, .55, 4.4, 8.8, .7, "Every primary external FPR gate fails.\nPerfect gold recall is not evidence of a safe operating point.", size=16, bold=True)
    LAYOUT.add_note(slide, "Denominators: internal P=14,326 / N=20,267; external gold P=69 / certified N=2,497.", y=5.32)
    titles.append(title)
    notes(slide, "All model thresholds remain validation-selected and unchanged. Counts and CP upper bounds are in the manuscript. Certified negatives cluster in 234 domains, so the exact binomial calculation does not remove dependence. Do not confuse gold-positive recall with precision or deployment utility.", ["secondary-metrics.csv", "primary-gates.csv"])
    slide = presentation.slides[4]
    title = "RQ1: structural gain; no incremental fixed-cascade gain"
    LAYOUT.set_title(slide, title)
    grid(slide, ["Contrast", "Recall gain (pp)", "Clustered 95% interval (pp)"], [["Internal L1 − length", "64.07", "[59.99, 67.63]"], ["Internal cascade − L1", "0.00", "[0.00, 0.00]"], ["Gold L1 − length", "76.81", "[66.67, 86.96]"], ["Gold cascade − L1", "0.00", "[0.00, 0.00]"]], widths=[.42,.22,.36], size=14, height=2.8)
    add_text(slide, .55, 4.3, 8.8, .75, "The fixed band selects 0/34,593 internal and 0/8,701 external rows.\nH1: not supported — 5 pass, 5 fail; all 10 measured.", size=16, bold=True)
    LAYOUT.add_note(slide, "2,000 domain-clustered replicates; internal positives: 9,757 domains; gold: 69 domains.", y=5.32)
    titles.append(title)
    notes(slide, "The band is unchanged and selects zero evaluation rows. The cascade therefore equals L1 on the evaluated streams; this does not show character representations are universally unnecessary. Strict-improvement lower bounds must exceed zero. Structural gains cannot cancel external specificity failures.", ["paired-contrasts.csv", "primary-results.json"])
    slide = presentation.slides[5]
    title = "RQ2: detecting departure did not make routing useful"
    panels(slide, title, "Monitoring", ["External alerts: 115/132 = 87.12%; passes ≥80%.", "Original audit: 28/252 = 11.11%; fails ≤5%.", "MMD and PSI: 132/132 external alerts; descriptive only."], "Routing consequences", ["Transformer routing: 8,253/8,701 = 94.85%.", "Gold recall gain: 0; interval [0,0].", "Certified FP: 2,261 → 2,479; policy FPR 99.2791%."], "H2: not supported — 1 pass, 3 fail. Overlapping windows do not establish independent harmful-drift events.")
    titles.append(title)
    notes(slide, "All 132 complete 256-row windows at stride 64 are the specified external departure population. Boundary -67.45792380813624 was not retuned. An alert affects the next 256 requests only; overlaps union and terminal activations truncate. These observed policy effects are not randomized causal remediation. MMD and PSI do not rescue H2.", ["external-monitor-windows.csv", "primary-results.json"])
    slide = presentation.slides[6]
    title = "RQ3: low transformer use did not guarantee low latency"
    panels(slide, title, "Primary operational gates", ["Reference physical calls: 0/10,000; passes ≤30%.", "Concurrency-64 pooled p95: 364.30 ms; fails ≤200 ms.", "Errors: 0/50,000; passes <0.1%."], "External safeguards", ["Cascade certified FPR: 90.55%; transformer: 99.28%.", "Both alert on 1,163/1,163 Tranco controls.", "Gold noninferiority passes, but does not establish specificity."], "H3: not supported — 3 pass, 5 fail. Primary cells were retained unchanged, not rerun for a better result.")
    titles.append(title)
    notes(slide, "Invocation uses only designated cell 1. Latency/errors use cells 21–25 and all 50,000 terminal latencies, including failures. Bounds use unrounded inputs; p95 is 364.3004100999999ms. Tranco is label-free alert rate. Noninferiority lower bound 0 ≥ -0.02 on 69 gold domains is a narrow passed component, not system viability.", ["operational-runs.csv", "primary-gates.csv"])
    slide = presentation.slides[7]
    title = "The complete service matrix includes adverse load results"
    LAYOUT.set_title(slide, title)
    grid(slide, ["Group", "Pooled p95 ms", "Errors / requests", "Physical calls"], [["Fixed 1%, c1", "1.49", "0 / 50,000", "0%"], ["Fixed 1%, c64", "364.30", "0 / 50,000", "0%"], ["Transformer, c64", "476.34", "74 / 50,000", "99.88%"], ["Transformer, c128", "970.86", "1,612 / 50,000", "96.964%"], ["Live shift, c1", "11.58", "0 / 43,505", "94.851%"]], widths=[.30,.22,.27,.21], size=13, height=3.0)
    add_text(slide, .55, 4.6, 8.8, .5, "All 25 groups and five repeats retained; full matrix in appendices.", size=16, bold=True)
    LAYOUT.add_note(slide, "Closed-loop throughput is not production capacity. Under-100% transformer attempts can reflect failures.", y=5.32)
    titles.append(title)
    notes(slide, "Fixed 1% c8 attempted throughput 1243.93–1272.10 requests/s versus 398.91–502.69 at c64. All three fixed c128 prevalence groups have errors: 120, 36, 59 respectively for 1%, 0.1%, 5%. The live shift uses a different stream, 8,701 requests/run and c1; it cannot substitute for the designated H3 benchmark. Across all five live runs there are 41,265 physical attempts/43,505 requests and exact offline/live trace agreement. Two sessions and fixed workload order confound session effects with some descriptive comparisons.", ["operational-groups.csv", "operational-runs.csv"])
    slide = presentation.slides[8]
    title = "Secondary evidence qualifies—not repairs—the primary result"
    panels(slide, title, "Completed coverage", ["153 population/model records; all seeds and comparators.", "43 mixed-class metric sets; 430 calibration bins; 129 prevalence projections.", "Full low-FPR curves, McNemar/Holm, MMD/PSI and label-free probes."], "Important qualifications", ["RF certified FPR: 91.35%; no secondary model is promoted.", "Source-only fit is unidentifiable; permutation audit limits remain explicit.", "Seed 42 runtime differs; probes do not establish correctness or robustness."], "No best-seed selection, post-test threshold replacement or fabricated negative-control success.")
    titles.append(title)
    notes(slide, "The frozen secondary schema provides full mixed-class metrics only internally and for gold+certified. Gold/certified count strata, positive-tier recall and Tranco label-free alerts remain separate. AP and ECE depend on this observed mixture, not deployment prevalence. Permutation consumed-label digests were not retained; varying ranking metrics cannot certify leakage or its absence. Original+3 separate probe streams each have 16,370 rows and 252 overlapping windows, with all score/decision/alert transitions exported.", ["secondary-verification.json", "complete-secondary-results.json"])
    slide = presentation.slides[9]
    title = "Technical value: connect prediction, routing and service"
    panels(slide, title, "What the evidence explains", ["Structural recall gains do not transfer to external specificity.", "Zero-band routing reveals a cascade with no incremental benefit.", "A shift alarm expands work without improving the target errors."], "What the work contributes", ["Joint feasibility analysis: 22 linked, unchanged checks.", "1,243,505 measured requests; all 125 cells and actual inference counters.", "Reusable dataflow, frozen evaluation, evidence map and auditable recovery."], "The value is a traceable systems explanation and reusable evaluation—not a claim to invent the component algorithms.")
    titles.append(title)
    notes(slide, "The dissertation now explains the implementation through data preparation, structural/character inference, monitoring, future-only routing and separate evaluation. Figure 3.1 and Table 3.1 map those planes to evidence. Sections 5.2.1–5.2.2 explain five technical contributions: joint feasibility, representation-to-operation analysis, monitoring-to-action analysis, physical HTTP measurement and auditable recovery. The complete matrix has 1,243,505 measured client requests excluding warmup and preserves all 1,901 errors; this overall count never replaces the designated 50,000-request H3 denominator. Appendix A maps every question to its tables and aggregate artifacts. The value is an inspectable systems explanation, not a universal impossibility theorem or first-ever claim. Component methods have prior art. A new design needs a separately specified future evaluation rather than retrospective repair of current gates.", ["primary-results.json", "operational-groups.csv", "operational-runs.csv"])
    slide = presentation.slides[10]
    title = "Completed dissertation: traceable from question to evidence"
    panels(slide, title, "Research and narrative", ["Exact RQ1–RQ3 answers; complete H1–H3 decisions.", "Architecture, technical contribution and evidence-navigation maps.", "All 22 checks, 125 cells, 25 groups and declared secondary evidence."], "Finished local deliverables", ["GWU-template manuscript: current abstract, contents, lists and references.", "Editable DOCX and directly rendered PDF; advisor deck and notes.", "Aggregate data, contracts, verification records and integrity manifest."], "Original manuscript and credential wording preserved. No institutional certification or advisor approval is asserted.")
    titles.append(title)
    notes(slide, "The current copy follows the supplied 2026 GWU D.Eng. Praxis template with a new abstract, page-linked contents, current tables/figure lists, Roman preliminary pagination and Arabic body numbering. It includes the five-chapter body and an evidence/reproduction appendix. The PDF is rendered directly from the DOCX, not from a differently paginated Markdown edition. The author authorized updating the new copy while preserving credentials and the original. Committee information is carried forward without asserting a passed examination or institutional approval. No advisor message is sent. Unsupported hypotheses are completed research results, and all adverse observations, methodological amendments and evidence limitations remain disclosed.", ["verification.json", "secondary-verification.json"])
    gates = records("primary-gates.csv")
    for index, hypothesis in enumerate(("H1", "H2", "H3"), start=11):
        slide = presentation.slides[index]
        LAYOUT.clear_slide_body(slide)
        title = f"Appendix: every {hypothesis} gate"
        LAYOUT.set_title(slide, title)
        rows = []
        for row in gates:
            if row["hypothesis"] != hypothesis:
                continue
            latency = row["name"] == "http_pooled_p95_ms"
            observed = f"{float(row['estimate']):.4f} ms" if latency else f"{100 * float(row['estimate']):.4f}%"
            threshold = f"{float(row['threshold']):g} ms" if latency else f"{100 * float(row['threshold']):g}%"
            if "_minus_" in row["name"]:
                observed = f"{100 * float(row['estimate']):.4f} pp"
                threshold = f"{100 * float(row['threshold']):g} pp"
            rows.append([row["name"].replace("_", " "), observed, row["operator"] + " " + threshold, row["status"].upper()])
        grid(slide, ["Gate", "Operand", "Required", "Result"], rows, widths=[.52,.18,.18,.12], size=11, height=3.9)
        LAYOUT.add_note(slide, "Contrast operands are lower confidence bounds; all decisions use full precision. No missing gates.", y=5.45, font_size=10.8)
        titles.append(title)
        notes(slide, "Complete gate list. Count denominators and CP bounds are in primary-gates.csv; point differences, bootstrap intervals and domain counts are in paired-contrasts.csv. Conjunctive support requires every component. Unsupported is not synonymous with all components failing.", ["primary-gates.csv", "paired-contrasts.csv"])
    groups = records("operational-groups.csv")
    for index, subset in enumerate((groups[:9], groups[9:18], groups[18:]), start=14):
        slide = presentation.slides[index]
        LAYOUT.clear_slide_body(slide)
        title = f"Appendix: full operational matrix ({index - 13}/3)"
        LAYOUT.set_title(slide, title)
        rows = []
        for group in subset:
            name = {"fixed_cascade": "Fixed", "transformer_only": "Transformer", "shift_period": "Shift"}[group["workload"]]
            prevalence = "external" if not group["prevalence_basis_points"] else f"{float(group['prevalence_basis_points']) / 100:g}%"
            rows.append([f"{name} {prevalence} c{group['concurrency']}", *[f"{float(group[key]):.2f}" for key in ("p50_ms", "p95_ms", "p99_ms")], f"{group['request_errors']}/{group['request_count']}"])
        grid(slide, ["Group (5 repeats)", "p50 ms", "p95 ms", "p99 ms", "Errors / requests"], rows, widths=[.35,.13,.13,.13,.26], size=11.7, height=3.9)
        LAYOUT.add_note(slide, "Warmup excluded; terminal latencies include failures. All individual throughput/counter records supplied.", y=5.45, font_size=10.8)
        titles.append(title)
        notes(slide, "These are pooled five-run values, not the best run or averages of quantiles. All ordinal membership, physical counters, elapsed/drain time, attempted throughput and successful throughput are in the run CSV. Group71–75 spans sessions. No post hoc adjustment separates workload from session effects.", ["operational-groups.csv", "operational-runs.csv"])
    properties = presentation.core_properties
    properties.title = "Praxis Advisor Results — October 1, 2026"
    for slide in list(presentation.slides)[1:]:
        title = slide.shapes.title
        title.left, title.top = Inches(.48), Inches(.22)
        title.width, title.height = Inches(9.02), Inches(.85)
    properties.subject = "Complete frozen evaluation: RQ1–RQ3 and H1–H3"
    properties.author = properties.last_modified_by = "Krti Tallam"
    properties.created = properties.modified = datetime.now(timezone.utc)
    for name in ("comments", "keywords", "category", "content_status", "identifier", "version"):
        setattr(properties, name, "")
    presentation.save(OUTPUT)
    temporary = OUTPUT.with_suffix(".rewrite.pptx")
    with ZipFile(OUTPUT) as source, ZipFile(temporary, "w", compression=ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "docProps/app.xml":
                root = etree.fromstring(data)
                for node in root.xpath("//*[local-name()='TitlesOfParts' or local-name()='HeadingPairs']"):
                    node.getparent().remove(node)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif entry.filename.startswith("ppt/") and entry.filename.endswith(".xml"):
                root = etree.fromstring(data)
                for node in root.xpath("//*[local-name()='fld' and starts-with(@type,'datetime')]/*[local-name()='t']"):
                    node.text = "10/1/26"
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(entry, data)
    temporary.replace(OUTPUT)
    if hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() != original_hash:
        raise RuntimeError("Historical deck changed")
    print(OUTPUT)
    print(f"{len(presentation.slides)} slides; historical SHA-256 unchanged: {original_hash}")


if __name__ == "__main__":
    main()
